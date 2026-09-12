# -*- coding: utf-8 -*-
"""
q3_produce.py -- 问题三生产计算: outputs/result3.xlsx, 表5, registry_q3.csv.

生产配置 (problem3.md §11.1): M=200 (dr=0.01 cm), rtol=1e-9,
atol=(1e-6 K, 1e-9 kg/kg), max_step=3600 s, tau=600 s, 解析稀疏 Jacobian.
另跑 M=100/400 作 Richardson 外推 (M=200/400 差 + 实测收敛阶, 与 q3_verify 独立).

输出:
  outputs/result3.xlsx            水分浓度, n60 x 21, 4 位小数 (Sheet1, 模板格式)
  outputs/table5_moisture.md/.csv 论文表5 (每 6 h + t_dry 末行)
  outputs/q3_production.log       运行记录
  outputs/registry_q3.csv         生产注册表 (论文对账用)

运行:  python src/q3_produce.py
"""

from __future__ import annotations

import csv
import os
import sys
import time

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from q3_solve import (PROD_ATOL_C, PROD_ATOL_T, PROD_MAX_STEP, PROD_M,      # noqa: E402
                      PROD_RTOL, run_parallel, write_registry)

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "outputs")

R_OUT_CM = np.round(np.arange(0.0, 2.0001, 0.1), 10)
TAB_COLS = (0, 5, 10, 15, 20)                     # 表5 的列: r=0,0.5,1,1.5,2 cm

LOG = []
REG = []


def say(s=""):
    print(s, flush=True)
    LOG.append(s)


def radd(id_, q, v, u, unc="", note=""):
    REG.append([id_, q, v, u, unc, "outputs/q3_production.log",
                "python src/q3_produce.py", note])


def main():
    import scipy
    t_all = time.time()
    say("=" * 96)
    say("问题三生产计算 (MOL + 自适应 BDF + 终止事件)")
    say(f"环境: python {sys.version.split()[0]}, scipy {scipy.__version__}, "
        f"numpy {np.__version__}")
    say(f"生产配置: M={PROD_M}, rtol={PROD_RTOL:g}, atol=({PROD_ATOL_T:g} K, "
        f"{PROD_ATOL_C:g} kg/kg), max_step={PROD_MAX_STEP:g} s, "
        f"tau=600 s, 解析稀疏 Jacobian")
    say(f"输出半径 (cm): {list(R_OUT_CM)}")
    say("=" * 96)

    res = run_parallel([dict(tag="prod", M=PROD_M, rtol=PROD_RTOL, sample=True,
                             crossings=True),
                        dict(tag="M100", M=100, rtol=1e-9),
                        dict(tag="M400", M=400, rtol=1e-9)], workers=3)
    prod, m100, m400 = res["prod"], res["M100"], res["M400"]
    t_dry = prod["t_dry"]

    # ------------------------------------------------------------------
    # 收敛性与 Richardson 外推 (生产口径)
    # ------------------------------------------------------------------
    d1 = prod["t_dry"] - m100["t_dry"]
    d2 = m400["t_dry"] - prod["t_dry"]
    p = float(np.log2(d1 / d2))
    rich = m400["t_dry"] + d2 / (2.0 ** p - 1.0)
    say("")
    say("网格收敛 (rtol=1e-9):")
    say(f"  M=100: {m100['t_dry']:.3f} s ({m100['t_dry'] / 3600.0:.5f} h)")
    say(f"  M=200: {prod['t_dry']:.3f} s ({prod['t_dry'] / 3600.0:.5f} h)  [生产]")
    say(f"  M=400: {m400['t_dry']:.3f} s ({m400['t_dry'] / 3600.0:.5f} h)")
    say(f"  d1={d1:.3f} s, d2={d2:.3f} s, p=log2(d1/d2)={p:.4f}")
    say(f"  Richardson 外推 t* = {rich:.1f} s = {rich / 3600.0:.4f} h")

    say("")
    say("=" * 96)
    say(f"生产解: t_dry = {t_dry:.3f} s = {t_dry / 3600.0:.4f} h "
        f"(约 {int(t_dry // 86400)} 天 {(t_dry % 86400) / 3600:.1f} 小时)")
    say(f"  步数 {prod['nsteps']}, nfev={prod['nfev']}, njev={prod['njev']}, "
        f"nlu={prod['nlu']}, 用时 {prod['wall']:.1f} s")
    say(f"  60 s 输出网格: n60 = {prod['n60']} 行 (末行 t = "
        f"{prod['t60'][-1]:.0f} s); 终态 max_i C_i = {prod['maxC_dry']:.12f}")
    say("=" * 96)

    # ------------------------------------------------------------------
    # 表 5
    # ------------------------------------------------------------------
    C60, t60 = prod["C60"], prod["t60"]
    t6 = 21600.0 * np.arange(1, int(np.floor(t_dry / 21600.0)) + 1)
    t6_h = t6 / 3600.0
    t_dry_h = t_dry / 3600.0
    say("")
    say("表5  药材烘干过程的水分浓度 (kg/kg)")
    say("  t/h      " + "".join(f"{d:>9}" for d in ("0", "0.5", "1", "1.5", "2")))
    rows5 = []
    for tt in t6:
        i = int(round(tt / 60.0)) - 1
        vals = [float(C60[i, j]) for j in TAB_COLS]
        rows5.append((f"{tt / 3600.0:.0f}", vals))
        say(f"  {tt / 3600.0:>5.0f}    " + "".join(f"{v:>9.4f}" for v in vals))
    vals_dry = [float(prod["C_dry"][j]) for j in TAB_COLS]
    rows5.append((f"{t_dry_h:.2f}", vals_dry))
    say(f"  {t_dry_h:>5.2f}    " + "".join(f"{v:>9.4f}" for v in vals_dry)
        + "   <- t_dry")

    for name, hdr in (("table5_moisture.md", True), ("table5_moisture.csv", False)):
        pth = os.path.join(OUT, name)
        if hdr:
            with open(pth, "w", encoding="utf-8") as f:
                f.write("| 时间/h | 0 | 0.5 | 1 | 1.5 | 2 |\n|---|---|---|---|---|---|\n")
                for lbl, vals in rows5:
                    f.write(f"| {lbl} | " + " | ".join(f"{v:.4f}" for v in vals)
                            + " |\n")
        else:
            with open(pth, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(["时间/h", "0", "0.5", "1", "1.5", "2"])
                for lbl, vals in rows5:
                    w.writerow([lbl] + [f"{v:.4f}" for v in vals])

    # ------------------------------------------------------------------
    # result3.xlsx
    # ------------------------------------------------------------------
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for j, rc in enumerate(R_OUT_CM):
        ws.cell(row=1, column=2 + j, value=float(rc))
    for i, tt in enumerate(t60):
        ws.cell(row=2 + i, column=1, value=int(round(tt)))
        for j in range(len(R_OUT_CM)):
            ws.cell(row=2 + i, column=2 + j, value=round(float(C60[i, j]), 4))
    pth = os.path.join(OUT, "result3.xlsx")
    wb.save(pth)
    say("")
    say(f"结果文件: {pth}  (Sheet1: {ws.max_row - 1} 行 x {ws.max_column - 1} 列, "
        f"A 列 60~{int(t60[-1])} s)")

    # ------------------------------------------------------------------
    # 达标次序与物理量摘要
    # ------------------------------------------------------------------
    cross = prod["crossings"]                     # {r_cm: t_s 或 None}
    say("")
    say("达标次序 (C 降到 0.15 kg/kg 的时刻, brentq 精化):")
    for rc in (0.0, 0.5, 1.0, 1.5, 2.0):
        t_ref = cross[rc]
        say(f"  r = {rc:.1f} cm: {t_ref / 3600.0:.3f} h" if t_ref is not None else
            f"  r = {rc:.1f} cm: 全程未低于 0.15")
    say(f"  终态: 中心 C = {prod['C_dry'][0]:.4f}, 表面 C = {prod['C_dry'][20]:.4f} "
        f"kg/kg (环境 Cbar = 0.0500, 表面传质阻力使其高于环境)")
    say(f"  终态: 中心 T = {prod['T_dry'][0] - 273.15:.4f} degC, 表面 T = "
        f"{prod['T_dry'][20] - 273.15:.4f} degC")
    loss = 100.0 * (1.0 - prod["W_dry"] / prod["W0"])
    budget = abs(prod["W_dry"] - prod["W0"]
                 + float(np.trapezoid(prod["flux_w"], prod["t_flux"])))
    resW = budget / abs(float(np.trapezoid(prod["flux_w"], prod["t_flux"])))
    dCmono = float(np.max(np.diff(prod["C60"], axis=0)))
    say(f"  总水量: W(0)={prod['W0']:.6e} -> W(t_dry)={prod['W_dry']:.6e} "
        f"(per 2 pi), 失水率 {loss:.4f}%, 收支相对残差 {resW:.3e}")
    say(f"  C(r,t) 时间单调性: 最大正增量 {dCmono:.3e} kg/kg")

    # ------------------------------------------------------------------
    # 注册表
    # ------------------------------------------------------------------
    from q1_solve import load_attachment1
    from q3_solve import Q3Env
    t1, T1C, C1 = load_attachment1()
    env0 = Q3Env(t1, T1C, C1)
    es = env0.stats
    say("")
    radd("E01_n", "附件1 末段 (10800~14400 s) 采样点数", es["n"], "-", "环境外推")
    radd("E02_T_mean", "附件1 末段环境温度均值", es["T_mean"], "degC", "环境外推")
    radd("E03_T_std", "附件1 末段环境温度标准差", es["T_std"], "degC", "环境外推")
    radd("E04_T_range", "附件1 末段环境温度极差", es["T_max"] - es["T_min"],
         "degC", "环境外推")
    radd("E05_C_mean", "附件1 末段环境水分浓度均值", es["C_mean"], "kg/kg", "环境外推")
    radd("E06_C_std", "附件1 末段环境水分浓度标准差", es["C_std"], "kg/kg", "环境外推")
    radd("E07_C_range", "附件1 末段环境水分浓度极差", es["C_max"] - es["C_min"],
         "kg/kg", "环境外推")
    say("")
    radd("P00_M", "生产网格数 M", PROD_M, "-", "网格设置")
    radd("P00_rtol", "生产相对容限 rtol", PROD_RTOL, "-", "容限设置")
    radd("P00_atol_T", "生产绝对容限 (温度)", PROD_ATOL_T, "K", "容限设置")
    radd("P00_atol_C", "生产绝对容限 (水分)", PROD_ATOL_C, "kg/kg", "容限设置")
    radd("P00_max_step", "生产步长上界", PROD_MAX_STEP, "s", "容限设置")
    radd("P00_tau", "环境外推线性过渡段长度", 600.0, "s", "环境外推")
    radd("P00_Tbar", "恒温平台环境温度 Tbar", 50.0, "degC", "环境外推 (附件1 末段均值)")
    radd("P00_Cbar", "恒温平台环境水分浓度 Cbar", 0.05, "kg/kg", "环境外推 (附件1 末段均值)")
    radd("P00_n60", "60 s 输出网格行数", int(prod["n60"]), "-", "输出格式")
    radd("P00_t_last", "60 s 输出网格末行时刻", float(prod["t60"][-1]), "s", "输出格式")
    radd("P01_tdry", "烘干所需时间 t_dry", t_dry, "s", "生产解")
    radd("P01_tdry_h", "烘干所需时间 t_dry", t_dry / 3600.0, "h", "生产解")
    radd("P01_tdry_rich", "Richardson 外推校正后的 t_dry", rich, "s", "M=100/200/400 外推")
    radd("P01_tdry_rich_h", "Richardson 外推校正后的 t_dry", rich / 3600.0, "h",
         "M=100/200/400 外推")
    radd("P02_M100", "M=100 的 t_dry", m100["t_dry"], "s", "网格收敛")
    radd("P02_M400", "M=400 的 t_dry", m400["t_dry"], "s", "网格收敛")
    radd("P02_d1", "相邻网格差 d1 (M100->200)", d1, "s", "网格收敛")
    radd("P02_d2", "相邻网格差 d2 (M200->400)", d2, "s", "网格收敛")
    radd("P02_order", "实测空间收敛阶 p = log2(d1/d2)", p, "-", "网格收敛")
    radd("P08_spatial_unc", "空间离散不确定度 (Richardson 与 M=400 之差)",
         (rich - m400["t_dry"]) / 3600.0, "h", "网格收敛")
    radd("P03_nsteps", "生产求解步数", int(prod["nsteps"]), "-", "求解统计")
    radd("P03_nfev", "生产求解 RHS 调用次数", int(prod["nfev"]), "-", "求解统计")
    radd("P03_njev", "生产求解 Jacobian 更新次数", int(prod["njev"]), "-", "求解统计")
    radd("P03_nlu", "生产求解 LU 分解次数", int(prod["nlu"]), "-", "求解统计")
    radd("P04_maxC_dry", "终态 max_i C_i 与 C* 之差", prod["maxC_dry"] - 0.15,
         "kg/kg", "事件一致性")
    radd("P04_argmax", "终态最大值位置 (节点号)", int(prod["argmax_dry"]), "-",
         "事件一致性 (0 = 轴心)")
    radd("P05_C_center", "t_dry 中心含水率", float(prod["C_dry"][0]), "kg/kg", "生产解")
    radd("P05_C_r0p5", "t_dry r=0.5cm 含水率", float(prod["C_dry"][5]), "kg/kg", "生产解")
    radd("P05_C_r1", "t_dry r=1cm 含水率", float(prod["C_dry"][10]), "kg/kg", "生产解")
    radd("P05_C_r1p5", "t_dry r=1.5cm 含水率", float(prod["C_dry"][15]), "kg/kg", "生产解")
    radd("P05_C_surface", "t_dry 表面含水率", float(prod["C_dry"][20]), "kg/kg", "生产解")
    radd("P05_Csurf_rel", "t_dry 表面含水率相对平台 Cbar 的偏差",
         100.0 * (float(prod["C_dry"][20]) / 0.05 - 1.0), "%", "生产解")
    radd("P05_T_center", "t_dry 中心温度", float(prod["T_dry"][0] - 273.15), "degC", "生产解")
    radd("P05_T_surface", "t_dry 表面温度", float(prod["T_dry"][20] - 273.15), "degC", "生产解")
    radd("P06_W0", "初始总水量 W(0) (per 2 pi)", prod["W0"], "m^2", "生产解")
    radd("P06_W_dry", "终态总水量 W(t_dry) (per 2 pi)", prod["W_dry"], "m^2", "生产解")
    radd("P06_loss", "全程失水率", loss, "%", "生产解")
    radd("P06_budget", "水量收支相对残差", resW, "-", "守恒检验")
    radd("P06_C_mono", "C 时间单调性最大正增量", dCmono, "kg/kg", "极值原理检验")
    for rc in (0.0, 0.5, 1.0, 1.5, 2.0):
        t_ref = cross[rc]
        if t_ref is not None:
            radd(f"P07_cross_r{rc:g}", f"r={rc:g} cm 的达标时刻 (brentq 精化)",
                 t_ref / 3600.0, "h", "达标次序")
    for ri, (lbl, vals) in enumerate(rows5):
        for cj, rc in zip(range(5), (0.0, 0.5, 1.0, 1.5, 2.0)):
            radd(f"T5_r{ri + 1}_c{cj + 1}", f"表5 t={lbl} h, r={rc:g} cm",
                 vals[cj], "kg/kg", "生产解")
    for tt in (600, 3600, 10800, 86400, 172800, int(round(t60[-1]))):
        i = int(round(tt / 60.0)) - 1
        if 0 <= i < prod["n60"]:
            for j, rc in ((0, 0.0), (10, 1.0), (20, 2.0)):
                radd(f"X_C_{tt}_{rc:g}", f"result3 水分 t={tt} s, r={rc:g} cm",
                     round(float(C60[i, j]), 4), "kg/kg", "result3.xlsx 抽查")

    rp = os.path.join(OUT, "registry_q3.csv")
    write_registry(rp, REG)
    with open(os.path.join(OUT, "q3_production.log"), "w", encoding="utf-8") as f:
        f.write("\n".join(LOG) + "\n")
    say(f"注册表: {rp} ({len(REG)} 行)")
    say(f"总用时 {time.time() - t_all:.1f} s")


if __name__ == "__main__":
    main()
