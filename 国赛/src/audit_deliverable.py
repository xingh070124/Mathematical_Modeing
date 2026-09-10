"""
AUDIT SCRIPT 9 -- end-to-end check of the actual deliverable.

  (1) shape / finiteness / monotonicity of outputs/result1.xlsx;
  (2) agreement with an INDEPENDENT computation (this audit's own solver, exact
      half-CV surface, face-averaged D) at the five reporting radii;
  (3) agreement of outputs/table1_temperature.*, table2_moisture.* with the xlsx.

Run:  python src/audit_deliverable.py
"""

from __future__ import annotations

import csv
import os
import sys

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from audit_common import (ALPHA, C0, HCONV, KCOND, R0, T0K, geom, make_env,  # noqa: E402
                          run)
from audit_mass import mass_run                         # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

hr = lambda s: (print("=" * 94), print(s), print("=" * 94))
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "outputs", "result1.xlsx")


def read_sheet(ws):
    hdr = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        rows.append([float(v) for v in r])
    return hdr, np.array(rows)


if __name__ == "__main__":
    hr("(1) outputs/result1.xlsx")
    wb = load_workbook(XLSX, data_only=True, read_only=True)
    print(f"  sheets: {wb.sheetnames}")
    sheets = {}
    for name in wb.sheetnames:
        hdr, A = read_sheet(wb[name])
        sheets[name] = (hdr, A)
        print(f"  [{name}] header[0:5]={hdr[:5]} ... header[-3:]={hdr[-3:]}")
        print(f"        shape={A.shape}  all finite={bool(np.isfinite(A).all())}  "
              f"t: {A[0,0]:.0f} .. {A[-1,0]:.0f} s  dt={np.unique(np.diff(A[:,1-1]))[:3]}")
    wb.close()

    tcol = sheets[sheets and list(sheets)[0]][1][:, 0]
    print(f"  time column: n={len(tcol)}  t[0]={tcol[0]}  t[-1]={tcol[-1]}  "
          f"uniform dt={np.allclose(np.diff(tcol), tcol[1]-tcol[0])}  dt={tcol[1]-tcol[0]}")

    Thi, Thd = sheets[list(sheets)[0]][1][:, 1:], sheets[list(sheets)[0]][0]
    x = np.array([float(v) for v in sheets[list(sheets)[0]][0][1:]])
    Thd = sheets[list(sheets)[0]][0]
    Thr = np.array([float(v) for v in Thd[1:]])
    print(f"  radii header (温度): {Thr[:5]} ... {Thr[-3:]}   count={len(Thr)}")

    for name in sheets:
        H, A = sheets[name]
        X = A[:, 1:]
        inc = bool(np.all(np.diff(X[-1]) >= -1e-12))
        dec = bool(np.all(np.diff(X[-1]) <= 1e-12))
        print(f"  [{name}] min={X.min():.6f} max={X.max():.6f}  "
              f"at t=1800 radially monotone increasing: {inc} / decreasing: {dec}  "
              f"time-monotone (col 0): {bool(np.all(np.diff(X[:,0]) >= -1e-12))}")
    print("  expected: temperature INCREASES with r (surface hotter), moisture DECREASES with r")

    hr("(2) independent recomputation (audit solver: exact half-CV surface)")
    fT, fC = make_env("pchip")
    xq = np.arange(21) * 0.001          # 0, 0.1, ..., 2.0 cm
    print(f"  {'source':>34} {'max|dT| [K]':>14} {'max|dC| [kg/kg]':>17} {'C at r=0/1cm/R':>24}")
    for M, dt in ((800, 0.125), (1600, 0.02), (1600, 0.125)):
        for df_, tag in (("face", "face-avg D (repaired)"), ("node", "nodal D (user)")):
            o = mass_run(M, dt, 1800.0, fT, fC, Dface=df_)
            Tm = np.interp(xq, o["r"], o["T"])
            Cm = np.interp(xq, o["r"], o["C"])
            dT = np.max(np.abs(Tm - sheets["温度"][1][-1, 1:] - 273.15))
            dC = np.max(np.abs(Cm - sheets["水分浓度"][1][-1, 1:]))
            print(f"  {'M=%d dt=%.3f %s' % (M, dt, tag):>34} {dT:>14.3e} {dC:>17.3e}   "
                  f"{Cm[0]:.5f} {Cm[10]:.5f} {Cm[-1]:.5f}")
    print("  a small max|dC| (≈1e-5) identifies which discretisation the deliverable used")

    hr("(3) table1/table2 files vs the xlsx")
    for f, sheet in (("table1_temperature.csv", "温度"), ("table2_moisture.csv", "水分浓度")):
        p = os.path.join(ROOT, "outputs", f)
        rows = list(csv.reader(open(p, encoding="utf-8-sig")))
        A = sheets[sheet][1]
        times = np.array([float(r[0]) for r in rows[1:]])
        vals = np.array([[float(v) for v in r[1:]] for r in rows[1:]])
        idx = [int(np.where(A[:, 0] == t)[0][0]) for t in times]
        ref = A[np.ix_(idx, [1, 6, 11, 16, 21])]      # 0, 0.5, 1, 1.5, 2 cm
        print(f"  {f}: rows={len(rows)-1} times={times.astype(int).tolist()} "
              f"cols={rows[0][1:]}")
        print(f"    max|table - xlsx[0,0.5,1,1.5,2 cm]| = {np.max(np.abs(vals-ref)):.3e}  "
              f"xlsx already 4-dec: {bool(np.allclose(ref, np.round(ref, 4)))}")
        print(f"    table row t=1800 : {vals[-1].tolist()}")
        print(f"    xlsx  row t=1800 : {ref[-1].tolist()}")
