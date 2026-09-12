# -*- coding: utf-8 -*-
"""Decisive test for C4 (production uncertainty) and the "4-decimal reporting is
safe" claim.

Runs the production configuration at dt/2 (M=1600, dt=1/128 s) and compares the
4-decimal values at the 30 table sample points (and the whole grid) against
outputs/result2.xlsx (M=1600, dt=1/64 s).

If the 4th decimal changes anywhere, the paper's criterion
    "max|dC| = 3.8467e-5 < 5e-5  =>  四位小数报告是安全的"
is refuted as a guarantee: it compares a *difference between two solutions* to a
half-ulp without accounting for the residual error of the finer of the two.

Outputs outputs/_audit_q2_c4.log and outputs/_audit_q2_c4_table.csv
"""
from __future__ import annotations

import csv
import os
import sys
import time

import numpy as np
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from q1_solve import Env, load_attachment1  # noqa: E402
from q2_solve import Par, march, out_indices  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "outputs")
M, DT = 1600, 1.0 / 128

say = []


def log(s=""):
    print(s, flush=True)
    say.append(str(s))


def main():
    t0 = time.time()
    t1, T1, C1 = load_attachment1()
    env = Env(t1, T1, C1, method="pchip")
    par = Par()
    idx = out_indices(M)
    o = march(M, DT, 10800.0, env, par, out_idx=idx, out_every_s=1.0)
    log(f"[1] 加密解完成 M={M} dt={DT:.8f} 步数=691200*2 "
        f"用时 {time.time()-t0:.1f} s")
    T = o["T_snap"]
    C = o["C_snap"]

    wb = openpyxl.load_workbook(os.path.join(OUT, "result2.xlsx"),
                                data_only=True, read_only=True)
    rh = np.array([float(v) for v in next(wb["温度"].iter_rows(values_only=True))[1:]])
    Tp = np.array([[float(x) for x in r[1:]]
                   for r in list(wb["温度"].iter_rows(values_only=True))[1:]])
    Cp = np.array([[float(x) for x in r[1:]]
                   for r in list(wb["水分浓度"].iter_rows(values_only=True))[1:]])
    wb.close()

    log()
    log("=" * 100)
    log("全网格最大偏差 (M=1600, dt=1/64 vs dt=1/128)")
    log("=" * 100)
    dT = np.max(np.abs(T - Tp))
    dC = np.max(np.abs(C - Cp))
    log(f"  max|dT| = {dT:.6e} K      (注册表 P01 = 1.3807690038448e-05)")
    log(f"  max|dC| = {dC:.6e} kg/kg  (注册表 P02 = 2.9838719628916e-05)")
    it = np.unravel_index(np.argmax(np.abs(C - Cp)), C.shape)
    log(f"  max|dC| 出现在 t={it[0]+1} s, r={rh[it[1]]:.1f} cm  "
        f"(dt=1/128 解 {C[it]:.7f}, 生产解 {Cp[it]:.7f})")

    log()
    log("=" * 100)
    log("表 4 (水分 / kg/kg): 四位小数逐位比较")
    log("=" * 100)
    cols = [0, 5, 10, 15, 20]
    rows = {}
    ndiff = 0
    for h in (0.5, 1.0, 1.5, 2.0, 2.5, 3.0):
        ts = int(round(h * 3600))
        a = np.round(C[ts - 1][cols], 4)
        b = np.round(Cp[ts - 1][cols], 4)
        line = "  ".join(f"{x:.4f}" for x in b)
        line2 = "  ".join(f"{x:.4f}{'*' if abs(x-y) > 1e-12 else ' '}"
                          for x, y in zip(b, a))
        rows[h] = (b, a)
        k = int(np.sum(np.abs(a - b) > 1e-12))
        ndiff += k
        log(f"  t={h:3.1f} h  生产: {line}")
        log(f"            dt/2 : {line2}   差异位数 {k}/5")
    log(f"  表4 30 个值中，dt 减半后 4 位小数发生变化的: {ndiff} 个")

    log()
    log("=" * 100)
    log("表 3 (温度 / degC): 四位小数逐位比较")
    log("=" * 100)
    ndiffT = 0
    for h in (0.5, 1.0, 1.5, 2.0, 2.5, 3.0):
        ts = int(round(h * 3600))
        a = np.round(T[ts - 1][cols], 4)
        b = np.round(Tp[ts - 1][cols], 4)
        k = int(np.sum(np.abs(a - b) > 1e-12))
        ndiffT += k
        log(f"  t={h:3.1f} h  {' '.join(f'{x:.4f}' for x in b)}"
            f"   | dt/2 差异 {k}/5")
    log(f"  表3 30 个值中，dt 减半后 4 位小数发生变化的: {ndiffT} 个")

    log()
    log("=" * 100)
    log("全表 (10800x21) 4 位小数的变化规模")
    log("=" * 100)
    for lbl, V, Vp in (("温度", T, Tp), ("水分", C, Cp)):
        a = np.round(V, 4)
        b = np.round(Vp, 4)
        n = int(np.sum(np.abs(a - b) > 1e-12))
        log(f"  {lbl}: {n} / {a.size} 个格点的第 4 位小数发生变化 "
            f"({100*n/a.size:.3f}%)")

    with open(os.path.join(OUT, "_audit_q2_c4_table.csv"), "w", newline="",
              encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["t_h", "r_cm", "C_prod_dt1_64_4dp", "C_dt1_128_4dp",
                    "changed"])
        for h in (0.5, 1.0, 1.5, 2.0, 2.5, 3.0):
            ts = int(round(h * 3600))
            for j in cols:
                w.writerow([h, rh[j], f"{Cp[ts-1][j]:.4f}",
                            f"{C[ts-1][j]:.4f}",
                            int(abs(round(C[ts-1][j], 4)
                                    - round(Cp[ts-1][j], 4)) > 1e-12)])
    with open(os.path.join(OUT, "_audit_q2_c4.log"), "w", encoding="utf-8") as f:
        f.write("\n".join(say) + "\n")
    log()
    log(f"总用时 {time.time()-t0:.1f} s")


if __name__ == "__main__":
    main()
