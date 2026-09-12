# -*- coding: utf-8 -*-
"""AUDIT q4 - probe 2: read 附件2.xlsx and the result4.xlsx template RAW."""
import openpyxl

P2 = r"A题\附件\附件2.xlsx"
P4 = r"A题\附件\附件3\result4.xlsx"

wb = openpyxl.load_workbook(P2, data_only=True)
print("=" * 70)
print("附件2.xlsx sheets:", wb.sheetnames)
ws = wb.worksheets[0]
print("dims:", ws.dimensions, "max_row", ws.max_row, "max_col", ws.max_column)
rows = list(ws.iter_rows(values_only=True))
print("header row:", rows[0])
print("first 4 data rows:", rows[1:5])
print("last 3 data rows:", rows[-3:])
data = [r for r in rows[1:] if r[0] is not None and r[1] is not None]
print("n data rows:", len(data))
import numpy as np
t = np.array([float(r[0]) for r in data])
R = np.array([float(r[1]) for r in data])
print("t[0..3]", t[:4], "t[-1]", t[-1], "dt unique:", np.unique(np.diff(t)))
print("R[0]", R[0], "R[-1]", R[-1], "min R", R.min(), "max R", R.max())
print("monotone nonincreasing:", bool(np.all(np.diff(R) <= 0)), "n_up:",
      int((np.diff(R) > 0).sum()))
print("R<=1.5 first t:", t[int(np.argmax(R <= 1.5))])
print("R<=1.2 first t:", t[int(np.argmax(R <= 1.2))])
print("R at 6h(21600s):", R[int(np.argmin(np.abs(t - 21600)))])
print("R at 4h(14400s):", R[int(np.argmin(np.abs(t - 14400)))])

print()
print("=" * 70)
wb4 = openpyxl.load_workbook(P4, data_only=True)
print("result4.xlsx template sheets:", wb4.sheetnames)
for name in wb4.sheetnames:
    w = wb4[name]
    print("--- sheet", name, "dims", w.dimensions, "max_row", w.max_row, "max_col", w.max_column)
    for i, r in enumerate(w.iter_rows(values_only=True)):
        print("   row", i, r)
        if i >= 4:
            break
