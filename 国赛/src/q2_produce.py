# -*- coding: utf-8 -*-
"""
q2_produce.py -- 问题二生产计算: outputs/result2.xlsx, 表3/表4, 不确定度注册表.

生产设置由 PROD_M / PROD_DT 给定 (依据 q2_verify.py 的收敛性结论选取).
另做两组独立加密校核, 给出生产解的离散不确定度, 并与四位小数阈值 5e-5 比较。

输出:
  outputs/result2.xlsx                     温度 / 水分浓度, 10800 x 21, 4 位小数
  outputs/table3_temperature.md/.csv       论文表3 (3 h 内每 0.5 h)
  outputs/table4_moisture.md/.csv          论文表4
  outputs/q2_production.log                运行记录
  outputs/registry_q2_production.csv       生产不确定度与关键值注册表

运行:  python src/q2_produce.py
"""

from __future__ import annotations

import csv
import os
import sys
import time

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from q1_solve import Env, load_attachment1                      # noqa: E402
from q2_solve import (Par, march, out_indices, props, Dfun,     # noqa: E402
                      hevap_of, H_CONV, HM, R, T0K, C0)

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "outputs")

# ---------------------------------------------------------------------------
# 生产设置 (依据 outputs/q2_verify.log 的 V1/V2 收敛性结论)
#   M 必须为 20 的倍数, 使输出半径 0,0.1,...,2.0 cm 恰好落在节点上
# ---------------------------------------------------------------------------
PROD_M = 1600
PROD_DT = 1.0 / 64.0           # 0.015625 s, 整除 1 s
CHK_M = 3200                   # 网格加密校核 (M x2)
CHK_DT = 1.0 / 128.0           # 时间步加密校核 (dt/2)
T_END = 10800.0
OUT_EVERY = 1.0

R_OUT_CM = np.round(np.arange(0.0, 2.0001, 0.1), 10)
T_TAB_S = [1800, 3600, 5400, 7200, 9000, 10800]
T_TAB_H = [0.5, 1.0, 1.5, 2.0, 2.5, 3.0]
TAB_COLS = [0, 5, 10, 15, 20]

LOG = []


def say(s=""):
    print(s, flush=True)
    LOG.append(s)


def _job(cfg):
    """子进程: 跑一组设置, 只回传输出半径上的结果."""
    from q1_solve import Env as _Env, load_attachment1 as _load
    t1, T1, C1 = _load()
    env = _Env(t1, T1, C1, method="pchip")
    M, dt, t_end = cfg["M"], cfg["dt"], cfg["t_end"]
    idx = out_indices(M)
    t0 = time.time()
    o = march(M, dt, t_end, env, Par(), out_idx=idx)
    return cfg["tag"], dict(T=o["T_snap"], C=o["C_snap"], t=o["t_snap"],
                            W=o["W"], E=o["E"], Fh=o["Fh"], Fw=o["Fw"],
                            stats=o["stats"], M=M, dt=dt, wall=time.time() - t0)


def run_all(cfgs, workers=3):
    from concurrent.futures import ProcessPoolExecutor
    out = {}
    t0 = time.time()
    with ProcessPoolExecutor(max_workers=workers) as ex:
        for tag, o in ex.map(_job, cfgs):
            out[tag] = o
            print(f"  [{tag}] M={o['M']}, dt={o['dt']:.10g} s, "
                  f"{o['stats']['nsteps']} 步, 迭代均值 "
                  f"{o['stats']['iters'].mean():.3f}, 用时 {o['wall']:.1f} s",
                  flush=True)
    print(f"  [三组算例并行, 共用时 {time.time()-t0:.1f} s]", flush=True)
    return out



if __name__ == "__main__":
    t1, T1, C1 = load_attachment1()
    env = Env(t1, T1, C1, method="pchip")

    say("=" * 96)
    say("问题二生产计算 (线性单元 + 集中质量矩阵 + 后向 Euler + 严格 Newton)")
    say("=" * 96)
    say(f"  输出半径 (cm): {list(R_OUT_CM)}")
    say(f"  输出时刻: 1..{int(T_END)} s, 共 {int(T_END)} 个")
    say(f"  物性: 附录3;  边界: 含蒸发吸热 q_evap = H_evap h_m (C_R - C_inf)")
    say(f"  H_evap(T) = 2.4346e6 - 2.391e3 (T - 28 degC) J/kg")
    say("")
    say(f"[A] 生产解 M={PROD_M}, dt={PROD_DT:.10g} s")
    say(f"[B] 校核1 时间步加密 (dt/2), 同网格 M={PROD_M}")
    say(f"[C] 校核2 网格加密 (M x2), 同时间步 dt={PROD_DT:.10g} s")
    say("")

    cfgs = [dict(tag="prod", M=PROD_M, dt=PROD_DT, t_end=T_END),
            dict(tag="chk-t", M=PROD_M, dt=CHK_DT, t_end=T_END),
            dict(tag="chk-h", M=CHK_M, dt=PROD_DT, t_end=T_END)]
    res = run_all(cfgs, workers=3)
    prod, chkt, chkh = res["prod"], res["chk-t"], res["chk-h"]
    wall_p, wall_t, wall_h = prod["wall"], chkt["wall"], chkh["wall"]

    Tp, Cp = prod["T"], prod["C"]
    dTt = float(np.max(np.abs(Tp - chkt["T"])))
    dCt = float(np.max(np.abs(Cp - chkt["C"])))
    dTh = float(np.max(np.abs(Tp - chkh["T"])))
    dCh = float(np.max(np.abs(Cp - chkh["C"])))

    say("")
    say("=" * 96)
    say("生产解的离散不确定度 (与两组独立加密解比较)")
    say("=" * 96)
    say(f"  时间步加密 (dt/2):  max|dT| = {dTt:.4e} K    max|dC| = {dCt:.4e} kg/kg")
    say(f"  网格加密  (M x2):   max|dT| = {dTh:.4e} K    max|dC| = {dCh:.4e} kg/kg")
    uT = max(dTt, dTh)
    uC = max(dCt, dCh)
    say(f"  取两者较大者作为不确定度:  T {uT:.4e} K,  C {uC:.4e} kg/kg")
    say(f"  四位小数阈值 (半 ulp) = 5e-5")
    say(f"  结论: 温度 {'低于' if uT < 5e-5 else '!! 高于'} 阈值; "
        f"水分 {'低于' if uC < 5e-5 else '!! 高于'} 阈值")

    say("")
    say("  分时段最大偏差 (网格加密 M x2, 用于说明启动瞬态):")
    win_rows = []
    say(f"  {'时间窗':>16} {'max|dT| [K]':>16} {'max|dC| [kg/kg]':>20}")
    for lo, hi, lbl in ((1, 60, "t=1..60 s"), (61, 600, "t=61..600 s"),
                        (601, 3600, "t=601..3600 s"), (3601, 10800, "t=3601..10800 s"),
                        (1, 10800, "t=1..10800 s")):
        sl = slice(lo - 1, hi)
        wT = float(np.max(np.abs((Tp - chkh["T"])[sl])))
        wC = float(np.max(np.abs((Cp - chkh["C"])[sl])))
        win_rows.append((lbl, wT, wC))
        say(f"  {lbl:>16} {wT:>16.4e} {wC:>20.4e}")

    # ---------------- result2.xlsx ----------------
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "温度"
    ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for j, rc in enumerate(R_OUT_CM):
        ws.cell(row=1, column=2 + j, value=float(rc))
    for i, tt in enumerate(prod["t"]):
        ws.cell(row=2 + i, column=1, value=int(round(tt)))
        for j in range(len(R_OUT_CM)):
            ws.cell(row=2 + i, column=2 + j, value=round(float(Tp[i, j]) - 273.15, 4))

    ws2 = wb.create_sheet("水分浓度")
    ws2.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for j, rc in enumerate(R_OUT_CM):
        ws2.cell(row=1, column=2 + j, value=float(rc))
    for i, tt in enumerate(prod["t"]):
        ws2.cell(row=2 + i, column=1, value=int(round(tt)))
        for j in range(len(R_OUT_CM)):
            ws2.cell(row=2 + i, column=2 + j, value=round(float(Cp[i, j]), 4))

    p = os.path.join(OUT, "result2.xlsx")
    wb.save(p)
    say("")
    say(f"结果文件: {p}  (温度 {ws.max_row-1}x{ws.max_column-1}, "
        f"水分浓度 {ws2.max_row-1}x{ws2.max_column-1})")

    # ---------------- 表3 / 表4 ----------------
    for name, arr, unit in (("table3_temperature.md", Tp - 273.15, "degC"),
                            ("table4_moisture.md", Cp, "kg/kg")):
        with open(os.path.join(OUT, name), "w", encoding="utf-8") as f:
            f.write("| 时间/h | 0 | 0.5 | 1 | 1.5 | 2 |\n|---|---|---|---|---|---|\n")
            for ih, tt in enumerate(T_TAB_S):
                i = int(tt) - 1
                f.write(f"| {T_TAB_H[ih]:.1f} | "
                        + " | ".join(f"{arr[i, j]:.4f}" for j in TAB_COLS) + " |\n")
    for name, arr in (("table3_temperature.csv", Tp - 273.15),
                      ("table4_moisture.csv", Cp)):
        with open(os.path.join(OUT, name), "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["时间/h", "0", "0.5", "1", "1.5", "2"])
            for ih, tt in enumerate(T_TAB_S):
                i = int(tt) - 1
                w.writerow([f"{T_TAB_H[ih]:.1f}"]
                           + [f"{arr[i, j]:.4f}" for j in TAB_COLS])

    say("")
    say("=" * 96)
    say("表3 温度 (degC) — 生产解 M=%d, dt=%.10g s" % (PROD_M, PROD_DT))
    say("=" * 96)
    say("  t/h    " + "".join(f"{d:>11}" for d in ["0", "0.5", "1", "1.5", "2"]))
    for ih, tt in enumerate(T_TAB_S):
        i = int(tt) - 1
        say(f"  {T_TAB_H[ih]:>4.1f}   "
            + "".join(f"{Tp[i, j]-273.15:>11.4f}" for j in TAB_COLS))
    say("")
    say("=" * 96)
    say("表4 水分浓度 (kg/kg) — 生产解 M=%d, dt=%.10g s" % (PROD_M, PROD_DT))
    say("=" * 96)
    say("  t/h    " + "".join(f"{d:>11}" for d in ["0", "0.5", "1", "1.5", "2"]))
    for ih, tt in enumerate(T_TAB_S):
        i = int(tt) - 1
        say(f"  {T_TAB_H[ih]:>4.1f}   "
            + "".join(f"{Cp[i, j]:>11.4f}" for j in TAB_COLS))

    # ---------------- 物理量摘要 ----------------
    say("")
    say("=" * 96)
    say("物理量摘要 (t=10800 s)")
    say("=" * 96)
    Tinf_end = env.T(T_END)
    Cinf_end = env.C(T_END)
    say(f"  环境: T_inf={Tinf_end-273.15:.4f} degC, C_inf={Cinf_end:.5f} kg/kg")
    say(f"  温度: 中心 {Tp[-1,0]-273.15:.4f} degC, 表面 {Tp[-1,-1]-273.15:.4f} degC")
    say(f"  水分: 中心 {Cp[-1,0]:.6f}, 表面 {Cp[-1,-1]:.6f} kg/kg")
    say(f"  水量: W(0)={prod['stats']['W0']:.6e}, W(10800)={prod['W'][-1]:.6e} "
        f"(per 2 pi, 单位轴长); 失水 {100*(1-prod['W'][-1]/prod['stats']['W0']):.4f}%")
    say(f"  水量收支相对残差 = "
        f"{abs((prod['W'][-1]-prod['stats']['W0'])+prod['Fw'][-1])/abs(prod['Fw'][-1]):.3e}")
    qh, qw = (H_CONV * R * (Tp[-1, -1] - Tinf_end)
              + hevap_of(Tp[-1, -1], Par())[0] * HM * R * (Cp[-1, -1] - Cinf_end)), \
        HM * R * (Cp[-1, -1] - Cinf_end)
    say(f"  t=10800 s 表面热流: 对流 {H_CONV*R*(Tp[-1,-1]-Tinf_end):.6f}, "
        f"蒸发 {hevap_of(Tp[-1,-1], Par())[0]*HM*R*(Cp[-1,-1]-Cinf_end):.6f} "
        f"(per 2 pi, W/m)")
    say(f"  Newton: 总步数 {prod['stats']['nsteps']}, "
        f"平均迭代 {prod['stats']['iters'].mean():.4f}, "
        f"最大 {prod['stats']['iters'].max()}")

    # ---------------- 注册表 ----------------
    reg = []

    def radd(id_, q, v, u, unc, note=""):
        vs = v if isinstance(v, str) else (
            f"{v:.14g}" if isinstance(v, (int, float, np.floating)) else str(v))
        reg.append([id_, q, vs, u, unc, "outputs/q2_production.log",
                    "python src/q2_produce.py", note])

    radd("P00_M", "生产网格数 M", PROD_M, "-", "网格设置")
    radd("P00_dt", "生产时间步 dt", PROD_DT, "s", "网格设置")
    radd("P00_nsteps", "生产总步数", int(prod["stats"]["nsteps"]), "-", "网格设置")
    radd("P01", "时间步加密 (dt/2): 温度最大偏差", dTt, "K", "dt/2 校核")
    radd("P02", "时间步加密 (dt/2): 水分最大偏差", dCt, "kg/kg", "dt/2 校核")
    radd("P03", "网格加密 (M x2): 温度最大偏差", dTh, "K", "M 加倍校核")
    radd("P04", "网格加密 (M x2): 水分最大偏差", dCh, "kg/kg", "M 加倍校核")
    radd("P05", "生产温度不确定度 (取较大者)", uT, "K", "离散不确定度")
    radd("P06", "生产水分不确定度 (取较大者)", uC, "kg/kg", "离散不确定度")
    radd("P07", "四位小数半 ulp 阈值", 5e-5, "-", "报告精度")
    for lbl, wT, wC in win_rows:
        tag = lbl.replace(" ", "").replace("=", "").replace("..", "_").replace(".", "p")
        radd(f"P10_{tag}_T", f"分时段网格加密最大偏差 (温度, {lbl})", wT, "K", "M 加倍校核")
        radd(f"P11_{tag}_C", f"分时段网格加密最大偏差 (水分, {lbl})", wC, "kg/kg", "M 加倍校核")
    radd("P20_Tinf_end", "t=10800 s 环境温度", Tinf_end - 273.15, "degC", "附件1")
    radd("P21_Cinf_end", "t=10800 s 环境含水率", Cinf_end, "kg/kg", "附件1")
    radd("P22_T_center", "t=10800 s 中心温度", Tp[-1, 0] - 273.15, "degC", "生产解")
    radd("P23_T_surface", "t=10800 s 表面温度", Tp[-1, -1] - 273.15, "degC", "生产解")
    radd("P24_C_center", "t=10800 s 中心含水率", Cp[-1, 0], "kg/kg", "生产解")
    radd("P25_C_surface", "t=10800 s 表面含水率", Cp[-1, -1], "kg/kg", "生产解")
    radd("P26_W0", "初始总水量 W(0) (per 2 pi)", prod["stats"]["W0"], "kg/kg*m^2", "生产解")
    radd("P27_W_end", "t=10800 s 总水量 W (per 2 pi)", prod["W"][-1], "kg/kg*m^2", "生产解")
    radd("P28_loss", "3 h 失水率", 100 * (1 - prod["W"][-1] / prod["stats"]["W0"]), "%", "生产解")
    radd("P29_Wres", "水量收支相对残差",
         abs((prod["W"][-1] - prod["stats"]["W0"]) + prod["Fw"][-1])
         / abs(prod["Fw"][-1]), "-", "离散守恒")
    radd("P30_it_mean", "Newton 平均迭代次数", float(prod["stats"]["iters"].mean()), "-", "迭代统计")
    radd("P31_it_max", "Newton 最大迭代次数", int(prod["stats"]["iters"].max()), "-", "迭代统计")
    radd("P32_wall", "生产运行墙钟时间", float(wall_p), "s", "机器相关, 不可复现")
    # 表3/表4 的每一个数值都进注册表 (供文稿引用)
    for ih, tt in enumerate(T_TAB_S):
        for j, rc in enumerate((0.0, 0.5, 1.0, 1.5, 2.0)):
            radd(f"T3_h{T_TAB_H[ih]:.1f}_r{rc:g}",
                 f"表3 温度 t={T_TAB_H[ih]:.1f} h, r={rc:g} cm",
                 float(Tp[int(tt) - 1, TAB_COLS[j]] - 273.15), "degC", "生产解")
            radd(f"T4_h{T_TAB_H[ih]:.1f}_r{rc:g}",
                 f"表4 水分 t={T_TAB_H[ih]:.1f} h, r={rc:g} cm",
                 float(Cp[int(tt) - 1, TAB_COLS[j]]), "kg/kg", "生产解")
    # result2.xlsx 的边界抽样 (用于对账抽查)
    for tt in (1, 60, 1800, 5400, 10800):
        for j, rc in enumerate((0.0, 1.0, 2.0)):
            radd(f"X_T_{tt}_{rc:g}", f"result2 温度 t={tt} s, r={rc:g} cm",
                 float(Tp[tt - 1, TAB_COLS[j * 2]] - 273.15), "degC", "生产解")
            radd(f"X_C_{tt}_{rc:g}", f"result2 水分 t={tt} s, r={rc:g} cm",
                 float(Cp[tt - 1, TAB_COLS[j * 2]]), "kg/kg", "生产解")

    rp = os.path.join(OUT, "registry_q2_production.csv")
    with open(rp, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "quantity", "value", "unit", "uncertainty", "source",
                    "command", "note"])
        w.writerows(reg)
    say("")
    say(f"注册表: {rp} ({len(reg)} 行)")

    with open(os.path.join(OUT, "q2_production.log"), "w", encoding="utf-8") as f:
        f.write("\n".join(LOG) + "\n")
    print(f"日志: {os.path.join(OUT, 'q2_production.log')}")
