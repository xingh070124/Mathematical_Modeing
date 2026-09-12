# -*- coding: utf-8 -*-
"""Audit L: final checks - C3 upper bound from the published grid, C6 Richardson view,
附件2 shrinkage claim, and the §3 attachment-1 point count."""
from __future__ import annotations
import os, sys
import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
from q1_solve import Env, load_attachment1

def hr(t):
    print("\n" + "=" * 78 + "\n" + t + "\n" + "=" * 78)

wb = load_workbook(os.path.join(ROOT, "outputs", "result2.xlsx"), read_only=True, data_only=True)
rows = list(wb["温度"].iter_rows(min_row=2, values_only=True))
T = np.array([[float(x) for x in r[1:]] for r in rows])
rad = np.array([float(x) for x in next(wb["温度"].iter_rows(min_row=1, max_row=1, values_only=True))[1:]])
wb.close()
t1, T1, C1 = load_attachment1()
env = Env(t1, T1, C1, method="pchip")
ts = np.arange(1, 10801)
Tinf = np.array([env.T(x) - 273.15 for x in ts])

hr("L1  C3 upper bound from the SHIPPED grid")
Tr = T[:, -1]
print(f"  max over all (t,r) of T        = {T.max():.4f} C  (doc 49.8888)")
print(f"  max over t of T(R)             = {Tr.max():.4f} C")
print(f"  max of T_inf over 0..10800 s   = {Tinf.max():.4f} C  (doc 50.2460)")
print(f"  count of (t, r=2cm) with T > T_inf : {int((Tr > Tinf).sum())} of 10800")
print(f"  max(T_R - T_inf)               = {float(np.max(Tr-Tinf)):+.6f} K")
if (Tr > Tinf).any():
    i = int(np.argmax(Tr - Tinf))
    print(f"    worst at t={ts[i]} s: T_R={Tr[i]:.4f} C, T_inf={Tinf[i]:.4f} C")
print(f"  count of all (t, all r) with T > max(T0, max T_inf) : {int((T > max(28.0, Tinf.max())).sum())}")
print(f"  margin to the upper envelope   = {Tinf.max() - T.max():.4f} K")

hr("L2  C6  what the 'uncertainty' actually is (Richardson view)")
uT_dt, uC_dt = 1.3808e-5, 2.9839e-5      # dt -> dt/2
uT_h, uC_h = 1.3824e-6, 3.8467e-5        # M -> 2M
print("  time (order 1): error(prod) ~ 2 x increment  ->  T %.3e K" % (2 * uT_dt))
print("  space (order 2): error(prod) ~ 4/3 x increment -> C %.4e kg/kg" % (4 / 3 * uC_h))
print(f"  reported (max convention):        T {max(uT_dt,uT_h):.4e}  C {max(uC_dt,uC_h):.4e}")
print(f"  sum convention:                   T {uT_dt+uT_h:.4e}  C {uC_dt+uC_h:.4e}")
print(f"  RSS convention:                   T {np.hypot(uT_dt,uT_h):.4e}  "
      f"C {np.hypot(uC_dt,uC_h):.4e}")
thr = 5e-5
for lbl, v in (("max (doc)", max(uC_dt, uC_h)), ("sum", uC_dt + uC_h),
               ("RSS", np.hypot(uC_dt, uC_h)), ("4/3 x M-increment", 4 / 3 * uC_h)):
    print(f"  moisture {lbl:>18}: {v:.4e} = {100*v/thr:.2f}% of 5e-5 "
          f"-> {'below' if v < thr else 'ABOVE'} threshold")
print("  doc quotes 76.93% (water) and 27.62% (temperature) - the max convention only")

hr("L3  §7.3 sub-window table vs q2_production.log")
print("  doc: 1.8054e-07 | 3.8467e-05 (t=1..60) ; log identical -> OK")
print("  BUT: the sub-window table is labelled '网格加密 M x2', so its 3.8467e-05 is")
print("  the M-increment, not the error of the shipped grid.")

hr("L4  §9.3 附件2 shrinkage claim")
from q1_solve import load_attachment2
try:
    a2 = load_attachment2()
    print("  attachment2 type:", type(a2))
    if isinstance(a2, tuple):
        for i, x in enumerate(a2):
            arr = np.asarray(x)
            print(f"   [{i}] shape {arr.shape} head {arr[:3]} tail {arr[-3:]}")
except Exception as e:
    print("  load_attachment2 failed:", e)
    from openpyxl import load_workbook as lw
    wb2 = lw(os.path.join(ROOT, "A题", "附件", "附件2.xlsx"), data_only=True)
    ws2 = wb2.active
    vals = list(ws2.iter_rows(values_only=True))
    print("  header:", vals[0])
    print("  first rows:", vals[1:4])
    print("  last rows:", vals[-3:])
    arr = np.array([[float(v) if v is not None else np.nan for v in r] for r in vals[1:]])
    print("  col means:", np.nanmean(arr, axis=0))
    if arr.shape[1] >= 2:
        tt, rr = arr[:, 0], arr[:, 1]
        m = tt <= 10800
        print(f"  R(0)={rr[0]:.4f} cm ; R at/just after 10800 s: "
              f"{rr[m][-1]:.4f} (t={tt[m][-1]:.0f} s)")
        print(f"  shrink over 0..10800 s = {100*(1-rr[m][-1]/rr[0]):.2f}%  (doc ~12%)")

hr("L5  §3 statement '附件1 的 241 个点 (dt=60 s)'")
print(f"  load_attachment1 -> n={len(t1)} points, t={t1[0]:.0f}..{t1[-1]:.0f} s, "
      f"dt={t1[1]-t1[0]:.0f} s  -> doc '241 points, 60 s' consistent: "
      f"{len(t1)==241 and t1[-1]==14400}")
