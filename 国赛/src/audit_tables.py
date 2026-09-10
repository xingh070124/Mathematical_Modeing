"""
AUDIT SCRIPT 10 -- final consolidated checks:

  (A) outputs/table1_temperature.csv vs outputs/table1_temperature.md vs
      outputs/result1.xlsx -- three artefacts that all claim to be 表1/表2.
  (B) the moisture surface convective coefficient ratio (claim C9 leading order).
  (C) the reconstruction of every claim number in the audit brief.

Run:  python src/audit_tables.py
"""

from __future__ import annotations

import csv
import os
import sys

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from audit_common import HCONV, HM, KCOND, R0, geom  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
hr = lambda s: (print("=" * 94), print(s), print("=" * 94))

COLS = [1, 6, 11, 16, 21]      # xlsx columns for 0, 0.5, 1, 1.5, 2 cm


def md_table(path):
    rows = []
    for ln in open(path, encoding="utf-8"):
        ln = ln.strip()
        if 2 <= ln.count("|") and not set(ln) <= set("|- "):
            cells = [c.strip() for c in ln.strip("|").split("|")]
            if cells[0] in ("时间/s",) or cells[0].startswith("---"):
                continue
            try:
                rows.append([float(c) for c in cells])
            except ValueError:
                pass
    return np.array(rows)


if __name__ == "__main__":
    hr("(A) three artefacts claiming to be 表1 / 表2")
    wb = load_workbook(os.path.join(ROOT, "outputs", "result1.xlsx"),
                       data_only=True, read_only=True)
    X = {}
    for name in wb.sheetnames:
        A = np.array([[float(v) for v in r] for r in wb[name].iter_rows(min_row=2, values_only=True)
                      if r[0] is not None])
        X[name] = A
    wb.close()

    for tag, sheet, csvf, mdf in (("表1 温度", "温度", "table1_temperature.csv", "table1_temperature.md"),
                                  ("表2 水分", "水分浓度", "table2_moisture.csv", "table2_moisture.md")):
        A = X[sheet]
        times = np.array([100, 300, 600, 900, 1200, 1500, 1800], dtype=float)
        idx = [int(np.where(A[:, 0] == t)[0][0]) for t in times]
        ref = A[np.ix_(idx, COLS)]
        rows = list(csv.reader(open(os.path.join(ROOT, "outputs", csvf), encoding="utf-8-sig")))
        csvv = np.array([[float(v) for v in r[1:]] for r in rows[1:]])
        mdv = md_table(os.path.join(ROOT, "outputs", mdf))[:, 1:]
        print(f"  {tag}: xlsx{ref.shape} csv{csvv.shape} md{mdv.shape}")
        print(f"    max|xlsx - md |  = {np.max(np.abs(ref - mdv)):.3e}")
        print(f"    max|xlsx - csv|  = {np.max(np.abs(ref - csvv)):.3e}")
        print(f"    max|csv  - md |  = {np.max(np.abs(csvv - mdv)):.3e}")
        d = np.abs(csvv - ref)
        bad = np.argwhere(d > 5e-5)
        print(f"    cells where |csv - xlsx| > 5e-5 (half a unit in the 4th decimal): {len(bad)}/35")
        for i, j in bad:
            print(f"      t={int(times[i])}s r={[0,0.5,1,1.5,2][j]}cm : csv={csvv[i,j]:.4f} "
                  f"md/xlsx={ref[i,j]:.4f}  diff={csvv[i,j]-ref[i,j]:+.4f}")

    hr("(B) moisture surface convective coefficient: user's 2*h_m/dr vs exact A_s*h_m/V_M")
    print(f"  {'M':>6} {'g_m_exact/(2 h_m/dr)':>22} {'relative deviation':>20}")
    for M in (50, 100, 200, 400, 800, 3200):
        dr, r, V, Am, Ap = geom(M)
        g_ex = 2 * np.pi * R0 * HM / V[M]
        g_us = 2 * HM / dr
        print(f"  {M:>6} {g_ex/g_us:>22.10f} {g_ex/g_us-1:>20.6e}")
    print(f"  closed form 1/(1-1/(4M)); identical to the thermal g_h factor")
    print(f"  thermal check M=200: g_h exact/user = "
          f"{(2*np.pi*R0*HCONV/(820.0*2600.0*geom(200)[2][200]))/(2*(KCOND/(820.0*2600.0))*HCONV/(KCOND*(R0/200))):.10f}")

    hr("(C) claim numbers, as recomputed in this audit")
    a = KCOND / (820.0 * 2600.0)
    print(f"  alpha                = {a:.10e} m^2/s   (user: 1.679e-7 ; rel +{a/1.679e-7-1:.6%})")
    dr200 = R0 / 200
    print(f"  M=200, alpha/dr^2    = {a/dr200**2:.6f} 1/s")
    print(f"  b_i correct          = 1/dt + {2*a/dr200**2:.6f}")
    print(f"  b_i user             = 1/dt + (i+0.5)/(i) * {a/dr200**2:.6f}")
    print(f"  row sum user (op)    = -(2i-1)/(2i dr^2) alpha ; i=1 -> {-a/(2*dr200**2):.6f}, "
          f"i=199 -> {-a*397/(398*dr200**2):.6f}")
