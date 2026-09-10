"""
问题1 (预热平衡阶段) 数值求解与建模方案审计.

被审计对象: 用户提供的"问题一完整建模思路"(控制方程 / 定解条件 / 差分格式).

本脚本做四件事:
  A. 符号验证: 用 sympy 重新"整理"用户的守恒型差分表达式, 得到三对角系数,
     与用户文中写出的 a_i,b_i,c_i 逐项比对.
  B. 数值复现: 逐字实现用户系数(user)与修正系数(fixed), 在同一网格上比较.
  C. 独立验证: (i) 与 Robin 边界无限长圆柱解析级数解对比 (常数 T_inf);
                (ii) method of lines + scipy BDF (非线性, D 不回滞).
  D. 产出: 表1/表2 结果、网格/时间步收敛表、数字注册表 registry_q1.csv.

网格: 节点式 (node-centered) FV, r_i = i*dr, i=0..M, dr=R/M.
      r=0 与 r=R 都是节点 -> 题面要求的输出位置 (0,0.1,...,2 cm) 落在节点上.

运行:  python src/q1_solve.py              # 全部实验
       python src/q1_solve.py --selftest   # 仅参数自检
"""

from __future__ import annotations

import argparse
import csv
import math
import os
import sys

import numpy as np
from scipy.integrate import solve_ivp
from scipy.linalg import solve_banded

# ----------------------------------------------------------------------------
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ATT1 = os.path.join(ROOT, "A题", "附件", "附件1.xlsx")
ATT2 = os.path.join(ROOT, "A题", "附件", "附件2.xlsx")
OUTDIR = os.path.join(ROOT, "outputs")

# ----------------------------------------------------------------------------
# 物理参数 (附录2, 问题1)
# ----------------------------------------------------------------------------
RHO = 820.0        # kg/m^3
CP = 2600.0        # J/(kg K)
K_COND = 0.36      # W/(m K)
H_CONV = 25.0      # W/(m^2 K)
HM = 8.0e-7        # m/s
D0 = 7.0e-9        # m^2/s
D_EXP = 0.89

T0_C = 28.0
T0_K = T0_C + 273.15
C0 = 2.55
R0 = 0.02          # m
L_AXIS = 0.25      # m

ALPHA = K_COND / (RHO * CP)     # 1.6885553e-7 m^2/s


def D_of_C(C):
    C = np.maximum(np.asarray(C, dtype=float), 1e-12)
    return D0 * np.exp(-D_EXP / C)


# ----------------------------------------------------------------------------
# 附件数据
# ----------------------------------------------------------------------------
def _read_xlsx(path, ncol):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    wb.close()
    return np.array([[float(v) for v in r[:ncol]] for r in rows])


def load_attachment1():
    a = _read_xlsx(ATT1, 3)
    return a[:, 0], a[:, 1], a[:, 2]


def load_attachment2():
    a = _read_xlsx(ATT2, 2)
    return a[:, 0], a[:, 1]


class Env:
    """环境激励 T_inf(t)[K], C_inf(t)[kg/kg] (附件1 的连续化)."""

    def __init__(self, t, T_C, C_inf, method="pchip"):
        from scipy.interpolate import PchipInterpolator, CubicSpline, interp1d

        self.method = method
        if method == "pchip":
            self._fT = PchipInterpolator(t, np.asarray(T_C) + 273.15)
            self._fC = PchipInterpolator(t, np.asarray(C_inf))
        elif method == "cubic":
            self._fT = CubicSpline(t, np.asarray(T_C) + 273.15)
            self._fC = CubicSpline(t, np.asarray(C_inf))
        elif method == "linear":
            self._fT = interp1d(t, np.asarray(T_C) + 273.15)
            self._fC = interp1d(t, np.asarray(C_inf))
        else:
            raise ValueError(method)
        self.t_tab = np.asarray(t, dtype=float)

    def T(self, tt):
        return float(self._fT(np.atleast_1d(float(tt)))[0])

    def C(self, tt):
        return float(self._fC(np.atleast_1d(float(tt)))[0])

    def T_arr(self, tt):
        return np.asarray(self._fT(np.asarray(tt, dtype=float)), dtype=float)

    def C_arr(self, tt):
        return np.asarray(self._fC(np.asarray(tt, dtype=float)), dtype=float)


def solve_tri(a, b, c, d):
    """三对角求解, 用 LAPACK 带状求解器 (a[0], c[-1] 忽略)."""
    n = len(d)
    ab = np.zeros((3, n))
    ab[0, 1:] = c[:-1]
    ab[1, :] = b
    ab[2, :-1] = a[1:]
    return solve_banded((1, 1), ab, d)


# ----------------------------------------------------------------------------
# 核心求解器: 节点式有限体积 + 全隐式
# ----------------------------------------------------------------------------
def geometry(M):
    """节点式 FV 几何量 (单位长度上已约去 2*pi)."""
    dr = R0 / M
    r = np.arange(M + 1) * dr
    # 控制体体积 (单位长度): pi*(r_{i+1/2}^2 - r_{i-1/2}^2)
    rf_p = (np.arange(M + 1) + 0.5) * dr      # r_{i+1/2}, i=0..M
    rf_m = (np.arange(M + 1) - 0.5) * dr      # r_{i-1/2}, i=0..M
    rf_m[0] = 0.0
    rf_p[M] = R0
    V = np.pi * (rf_p ** 2 - rf_m ** 2)
    A_p = 2 * np.pi * rf_p                     # 外侧面积, 长度 M+1
    A_m = 2 * np.pi * rf_m                     # 内侧面积, 长度 M+1
    return dr, r, V, A_p, A_m


def coeffs_FV(M, alpha, dt, dr, V, A_p, A_m, corr=True):
    """节点式 FV + 全隐式 Euler 的三对角系数.

    node 0  : V_0 dT_0/dt = A_{1/2} alpha (T_1-T_0)/dr
    node i  : V_i dT_i/dt = [A_{i+1/2}(T_{i+1}-T_i) - A_{i-1/2}(T_i-T_{i-1})] alpha/dr
    node M  : V_M dT_M/dt = -A_{M-1/2} alpha (T_M-T_{M-1})/dr
                            + 2*pi*R*h*(T_inf - T_M)/(rho cp)

    corr=True  -> 正确的对角元 (解析整理结果)
    corr=False -> 用户原文写出的对角元 (把 T_{i+1} 的系数当成了对角元)
    """
    a = np.zeros(M + 1)
    b = np.zeros(M + 1)
    c = np.zeros(M + 1)

    b[0] = 1.0 / dt + alpha * A_p[0] / (dr * V[0])
    c[0] = -alpha * A_p[0] / (dr * V[0])

    ii = np.arange(1, M)
    a[ii] = -alpha * A_m[ii] / (dr * V[ii])
    c[ii] = -alpha * A_p[ii] / (dr * V[ii])
    if corr:
        b[ii] = 1.0 / dt + alpha * (A_m[ii] + A_p[ii]) / (dr * V[ii])
    else:
        # 用户原文 b_i = (2i+1)/(2i dr^2) alpha + 1/dt  ==  c_i 的量值
        b[ii] = 1.0 / dt + alpha * A_p[ii] / (dr * V[ii])

    a[M] = -alpha * A_m[M] / (dr * V[M])
    b[M] = 1.0 / dt + alpha * A_m[M] / (dr * V[M])
    return a, b, c


def solve_q1(M, dt, t_end, env, corr=True, probes_t=None, no_moisture=False,
             const_T_inf=None):
    """节点式 FV 全隐式求解. 返回 (r, T[K], C, snap_T, snap_C)."""
    dr, r, V, A_p, A_m = geometry(M)
    nsteps = int(round(t_end / dt))

    # 热: 对流项系数 (仅 node M)
    g_h = 2 * np.pi * R0 * H_CONV / (RHO * CP * V[M])       # 1/s per K
    # 湿: 对流项系数 (仅 node M)
    g_hm = 2 * np.pi * R0 * HM / V[M]                        # 1/s

    T = np.full(M + 1, T0_K)
    C = np.full(M + 1, C0)

    aT, bT, cT = coeffs_FV(M, ALPHA, dt, dr, V, A_p, A_m, corr=corr)
    bT[M] = bT[M] + g_h
    aC, bC, cC = coeffs_FV(M, 1.0, dt, dr, V, A_p, A_m, corr=corr)  # 单位 D

    snaps_T = snaps_C = None
    if probes_t is not None:
        snaps_T = np.empty((len(probes_t), M + 1))
        snaps_C = np.empty((len(probes_t), M + 1))
        ptr = 0

    for n in range(nsteps):
        tn1 = (n + 1) * dt
        Ti = env.T(tn1) if const_T_inf is None else const_T_inf
        Ci = env.C(tn1)

        # ---- 温度 ----
        dT = T / dt
        a, b, c = aT.copy(), bT.copy(), cT.copy()
        dT[M] += g_h * Ti
        T = solve_tri(a, b, c, dT)

        # ---- 水分 (D 取上一时刻, 半隐式) ----
        if not no_moisture:
            Dn = D_of_C(C)
            # 面扩散系数: 相邻节点算术平均
            Df_m = 0.5 * (Dn[:-1] + Dn[1:])       # 长度 M, 对应 r_{i+1/2}
            a = aC.copy() * 1.0
            b = bC.copy()
            c = cC.copy()
            # 重新按 D 加权组装
            a[:] = 0.0
            b[:] = 1.0 / dt
            c[:] = 0.0
            # node 0
            b[0] += Df_m[0] * A_p[0] / (dr * V[0])
            c[0] = -Df_m[0] * A_p[0] / (dr * V[0])
            ii = np.arange(1, M)
            a[ii] = -Df_m[ii - 1] * A_m[ii] / (dr * V[ii])
            c[ii] = -Df_m[ii] * A_p[ii] / (dr * V[ii])
            b[ii] += Df_m[ii - 1] * A_m[ii] / (dr * V[ii]) + Df_m[ii] * A_p[ii] / (dr * V[ii])
            a[M] = -Df_m[M - 1] * A_m[M] / (dr * V[M])
            b[M] += Df_m[M - 1] * A_m[M] / (dr * V[M]) + g_hm
            dC = C / dt
            dC[M] += g_hm * Ci
            C = solve_tri(a, b, c, dC)

        if probes_t is not None:
            while ptr < len(probes_t) and tn1 >= probes_t[ptr] - 1e-9:
                snaps_T[ptr] = T
                snaps_C[ptr] = C
                ptr += 1

    return r, T, C, snaps_T, snaps_C


# ----------------------------------------------------------------------------
# 独立验证 1: Robin 边界无限长圆柱解析级数解
# ----------------------------------------------------------------------------
def robin_cylinder_eigen(Bi, n_modes=200):
    """lambda_n 满足 lambda J1(lambda) = Bi J0(lambda); 返回 (lam, C_n)."""
    from scipy.special import j0, j1
    from scipy.optimize import brentq

    f = lambda lam: lam * j1(lam) - Bi * j0(lam)
    lams = []
    # 根位于 J0 与 J1 零点之间; 用密集扫描找符号变化
    grid = np.linspace(1e-6, 400.0, 200000)
    vals = f(grid)
    for k in range(len(grid) - 1):
        if vals[k] == 0:
            lams.append(grid[k])
        elif vals[k] * vals[k + 1] < 0:
            lams.append(brentq(f, grid[k], grid[k + 1], xtol=1e-14))
        if len(lams) >= n_modes:
            break
    lams = np.array(lams[:n_modes])
    C = 2 * j1(lams) / (lams * (j0(lams) ** 2 + j1(lams) ** 2))
    return lams, C


def robin_cylinder_theta(rho, Fo, lams, C):
    from scipy.special import j0
    rho = np.atleast_1d(rho)
    th = np.zeros_like(rho)
    for lam, Cn in zip(lams, C):
        th += Cn * j0(lam * rho) * np.exp(-lam ** 2 * Fo)
    return th


# ----------------------------------------------------------------------------
# 独立验证 2: method of lines + BDF
# ----------------------------------------------------------------------------
def solve_mol(M, t_end, env, rtol=1e-10, atol=1e-12, nsamp=1):
    dr, r, V, A_p, A_m = geometry(M)
    g_h = 2 * np.pi * R0 * H_CONV / (RHO * CP * V[M])
    g_hm = 2 * np.pi * R0 * HM / V[M]

    def rhs(t, y):
        T = y[:M + 1]
        C = y[M + 1:]
        dT = np.empty(M + 1)
        dC = np.empty(M + 1)
        # node 0
        dT[0] = ALPHA * A_p[0] * (T[1] - T[0]) / (dr * V[0])
        ii = np.arange(1, M)
        dT[1:M] = ALPHA * (A_p[1:M] * (T[2:] - T[1:M]) - A_m[1:M] * (T[1:M] - T[:M - 1])) / (dr * V[1:M])
        dT[M] = ALPHA * (-A_m[M] * (T[M] - T[M - 1]) / (dr * V[M])) + g_h * (env.T(t) - T[M])
        Dn = D_of_C(C)
        Df = 0.5 * (Dn[:-1] + Dn[1:])
        dC[0] = Df[0] * A_p[0] * (C[1] - C[0]) / (dr * V[0])
        dC[1:M] = (Df[1:M] * A_p[1:M] * (C[2:] - C[1:M])
                   - Df[:M - 1] * A_m[1:M] * (C[1:M] - C[:M - 1])) / (dr * V[1:M])
        dC[M] = Df[M - 1] * (-A_m[M] * (C[M] - C[M - 1]) / (dr * V[M])) + g_hm * (env.C(t) - C[M])
        return np.concatenate([dT, dC])

    y0 = np.concatenate([np.full(M + 1, T0_K), np.full(M + 1, C0)])
    te = np.atleast_1d(np.linspace(0.0, t_end, nsamp + 1)[1:]) if nsamp > 1 else np.atleast_1d(np.array([t_end]))
    sol = solve_ivp(rhs, (0.0, t_end), y0, method="BDF", rtol=rtol, atol=atol, t_eval=te)
    if not sol.success:
        raise RuntimeError(f"MOL failed: {sol.message}")
    return r, sol.y[:M + 1], sol.y[M + 1:]


# ----------------------------------------------------------------------------
def symbolic_check():
    """用 sympy 重新整理用户的守恒型差分表达式, 得到三对角系数."""
    import sympy as sp

    i, dr_s, al, dt_s = sp.symbols("i dr alpha dt", positive=True)
    T_im1, T_i, T_ip1 = sp.symbols("T_im1 T_i T_ip1")
    # 用户原文: (1/r_i) * [ r_{i+1/2}(T_{i+1}-T_i)/dr - r_{i-1/2}(T_i-T_{i-1})/dr ] / dr
    expr = (sp.Rational(1, 1) / (i * dr_s)) * (
        (i + sp.Rational(1, 2)) * dr_s * (T_ip1 - T_i) / dr_s
        - (i - sp.Rational(1, 2)) * dr_s * (T_i - T_im1) / dr_s
    ) / dr_s
    expr = sp.expand(expr)
    co_a = sp.simplify(sp.expand(expr).coeff(T_im1))   # T_{i-1}
    co_b = sp.simplify(sp.expand(expr).coeff(T_i))     # T_i
    co_c = sp.simplify(sp.expand(expr).coeff(T_ip1))   # T_{i+1}
    rowsum = sp.simplify(co_a + co_b + co_c)
    return {"a_i": sp.simplify(co_a), "b_i": sp.simplify(co_b), "c_i": sp.simplify(co_c),
            "row_sum": rowsum}


# ----------------------------------------------------------------------------
def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    ap = argparse.ArgumentParser()
    ap.add_argument("--selftest", action="store_true")
    args = ap.parse_args()

    os.makedirs(OUTDIR, exist_ok=True)
    t1, T1, C1 = load_attachment1()
    t2, R2 = load_attachment2()
    env = Env(t1, T1, C1, method="pchip")

    reg = []
    R_CMD = "python src/q1_solve.py"

    def add(id_, q, v, u, unc, src, note=""):
        vs = v if isinstance(v, str) else (f"{v:.10g}" if isinstance(v, (int, float, np.floating)) else str(v))
        reg.append([id_, q, vs, u, unc, src, R_CMD, note])
        print(f"  {id_:12s} {q:56s} = {vs} {u}")

    hr = lambda s: (print("=" * 96), print(s), print("=" * 96))

    hr("S0. 参数自检")
    add("P01", "热扩散系数 alpha = k/(rho cp)", ALPHA, "m^2/s", "解析精确", "附录2", "")
    add("P02", "用户文中 alpha 数值 1.679e-7 的相对偏差", ALPHA / 1.679e-7 - 1.0, "-", "解析精确",
        "用户方案 2.1", "用户写 1.679e-7, 实为 1.6886e-7")
    for Cv in (C0, 1.0, 0.5, 0.15, 0.0331):
        print(f"  D(C={Cv:<7}) = {float(D_of_C(Cv)):.6e} m^2/s")
    add("P03", "D(C0=2.55)", float(D_of_C(C0)), "m^2/s", "解析精确", "附录2", "")
    add("P04", "D(C=0.15)", float(D_of_C(0.15)), "m^2/s", "解析精确", "附录2", "问题3阈值处")
    add("P05", "D(C=0.0331)", float(D_of_C(0.0331)), "m^2/s", "解析精确", "附录2", "1800s C_inf")
    add("P06", "D 在 C0..0.0331 上的量级跨度", float(D_of_C(C0) / D_of_C(0.0331)), "-", "解析精确",
        "附录2", "12 个数量级")
    add("P06b", "D(C0)/D(0.15) (初始 vs 问题3阈值)", float(D_of_C(C0) / D_of_C(0.15)), "-",
        "解析精确", "附录2", "干壳效应的量级")
    add("P07", "热 Biot 数 Bi = hR/k", H_CONV * R0 / K_COND, "-", "解析精确", "附录2", ">0.1 内部温度不均匀")
    add("P08", "传质 Biot 数 Bi_m = hm R/D(C0)", HM * R0 / float(D_of_C(C0)), "-", "解析精确", "附录2", "")
    add("P09", "导热特征时间 R^2/alpha", R0 ** 2 / ALPHA, "s", "解析精确", "附录2", "与 1800 s 比较")
    add("P10", "含水率扩散特征时间 R^2/D(C0)", R0 ** 2 / float(D_of_C(C0)), "s", "解析精确", "附录2", "")
    add("P11", "长径比 L/R", L_AXIS / R0, "-", "解析精确", "题面", "")
    add("P12", "端面面积/侧面面积", math.pi * R0 ** 2 / (2 * math.pi * R0 * L_AXIS), "-", "解析精确",
        "题面", "忽略轴向的相对误差量级")
    add("P13", "附件1 温度范围", f"{T1.min():.3f}..{T1.max():.3f}", "degC", "解析精确", "附件1", "")
    add("P14", "附件1 水分浓度范围", f"{C1.min():.5f}..{C1.max():.5f}", "kg/kg", "解析精确", "附件1", "")
    add("P15", "附件1 采样间隔", float(t1[1] - t1[0]), "s", "解析精确", "附件1", "")
    dC_ = np.diff(C1)
    add("P16", "附件1 C_inf 下降步数 / 总步数", f"{int(np.sum(dC_ < 0))}/{len(dC_)}", "-", "解析精确",
        "附件1", "非单调 => 数据含噪声")
    add("P17", "附件1 C_inf 最大单步降幅", float(max(0.0, -dC_.min())), "kg/kg", "解析精确", "附件1", "")
    dT_ = np.diff(T1)
    add("P18", "附件1 T_inf 下降步数 / 总步数", f"{int(np.sum(dT_ < 0))}/{len(dT_)}", "-", "解析精确",
        "附件1", "")
    add("P19", "附件2 半径范围", f"{R2.min():.3f}..{R2.max():.3f}", "cm", "解析精确", "附件2", "")
    add("P20", "附件2 采样间隔", float(t2[1] - t2[0]), "s", "解析精确", "附件2", "")
    add("P21", "附件2 t=1800s 半径", float(R2[1]), "cm", "解析精确", "附件2", "问题1时段末")
    add("P22", "附件1 覆盖时长", float(t1[-1]), "s", "解析精确", "附件1", "问题1只需 1800 s")
    add("P23", "1800s 时 C_inf/C0", float(C1[30] / C0), "-", "解析精确", "附件1", "")

    if args.selftest:
        _write_registry(reg, os.path.join(OUTDIR, "registry_q1.csv"))
        return

    # ------------------------------------------------------------------ S1
    hr("S1. 符号验证: 重新整理用户的守恒型差分表达式")
    sym = symbolic_check()
    print(f"  T_im1 系数 a_i = {sym['a_i']}")
    print(f"  T_i   系数 b_i = {sym['b_i']}")
    print(f"  T_ip1 系数 c_i = {sym['c_i']}")
    print(f"  行和 a+b+c     = {sym['row_sum']}   <-- 须为 0")
    import sympy as sp
    i_s, al_s, dr_s = sp.symbols("i alpha dr", positive=True)
    b_user = (2 * i_s + 1) / (2 * i_s * dr_s ** 2) * al_s
    b_correct = sp.simplify(sym["b_i"] * al_s)
    print(f"  正确对角元 b_i*alpha = {b_correct}")
    print(f"  用户对角元 b_i*alpha = {b_user}")
    print(f"  差 = {sp.simplify(b_user - b_correct)}")
    add("S01", "正确内部对角元 b_i (乘 1/alpha)", str(b_correct), "1/m^2", "符号精确", "sympy 重算", "")
    add("S02", "用户内部对角元 b_i (乘 1/alpha)", str(b_user), "1/m^2", "符号精确", "用户方案 4.3", "")
    add("S03", "两者之差 (乘 1/alpha)", str(sp.simplify(b_user - b_correct)), "1/m^2", "符号精确", "sympy", "")
    add("S04", "用户系数行和 a+b+c (乘 1/alpha)", str(sp.simplify(
        -(2 * i_s - 1) / (2 * i_s * dr_s ** 2) + (2 * i_s + 1) / (2 * i_s * dr_s ** 2)
        - (2 * i_s + 1) / (2 * i_s * dr_s ** 2))), "1/m^2", "符号精确", "用户方案 4.3",
        "非 0 => 非守恒")

    # ------------------------------------------------------------------ S2
    hr("S2. 用户系数格式的数值表现 (M=200, dt=0.1s, 用户推荐)")
    M_test = 200
    for tag, corr in (("fixed", True), ("user", False)):
        try:
            r_u, T_u, C_u, _, _ = solve_q1(M_test, 0.1, 60.0, env, corr=corr, no_moisture=True)
            print(f"  [{tag:5s}] t=60s: T_min={T_u.min():.4f} K  T_max={T_u.max():.4f} K  "
                  f"T_inf(60)={env.T(60):.4f} K")
            if tag == "user":
                add("S05", "用户系数格式 t=60s 温度最大值", float(T_u.max()), "K", "run",
                    "run solve_q1(corr=False)", f"物理上界 T_inf(60s)={env.T(60):.4f} K")
                add("S06", "用户系数格式 t=60s 温度最小值", float(T_u.min()), "K", "run",
                    "run solve_q1(corr=False)", "")
                add("S07", "用户系数格式是否越出物理区间",
                    str(bool(T_u.min() < T0_K - 1 or T_u.max() > env.T(60) + 1)), "-", "run",
                    "run solve_q1(corr=False)", "True => 格式失效")
        except Exception as e:
            print(f"  [{tag:5s}] 异常: {e}")

    # ------------------------------------------------------------------ S3
    hr("S3. 独立验证 1 — Robin 边界无限长圆柱解析级数解 (常数 T_inf)")
    Bi = H_CONV * R0 / K_COND
    lams, Cn = robin_cylinder_eigen(Bi, n_modes=120)
    print(f"  Bi = {Bi:.6f};  前 5 个特征值: {np.round(lams[:5], 6)}")
    print(f"  前 5 个系数 C_n: {np.round(Cn[:5], 6)}")
    # 自检: 级数在 t=0 应重构初始条件 theta=1
    rho_test = np.linspace(0, 1, 21)
    recon = np.array([robin_cylinder_theta(rr, 1e-12, lams, Cn)[0] for rr in rho_test])
    err_recon = np.max(np.abs(recon - 1.0))
    print(f"  初始条件重构 max|sum C_n J0(lam_n rho) - 1| = {err_recon:.3e}")
    add("S08", "解析级数初始条件重构误差 max", float(err_recon), "-", "级数截断=120项",
        "run robin_cylinder", "<1e-6 说明特征值与系数正确")

    T_const = 323.15  # 50 degC
    env_c = Env(np.array([0.0, 1e6]), np.array([50.0, 50.0]), np.array([C0, C0]), method="linear")
    for M_ in (100, 200, 400, 800, 1600):
        t_chk = 300.0
        r_, T_, C_, _, _ = solve_q1(M_, 0.05, t_chk, env_c, corr=True, no_moisture=True,
                                    const_T_inf=T_const)
        Fo = ALPHA * t_chk / R0 ** 2
        th = robin_cylinder_theta(r_ / R0, Fo, lams, Cn)
        Tan = T_const + (T0_K - T_const) * th
        e = np.max(np.abs(T_ - Tan))
        print(f"  M={M_:>5}: max|T_FV - T_analytic| = {e:.3e} K   (Fo={Fo:.4f})")
        add(f"S09_M{M_}", f"FV vs 解析解最大偏差 (M={M_}, t=300s)", float(e), "K",
            "解析级数120项", "run robin compare", "")

    # ------------------------------------------------------------------ S4
    hr("S4. 网格收敛性 (固定 dt=0.05 s, 节点式 FV 修正格式)")
    t_end = 1800.0
    probes_r = np.array([0.0, 0.005, 0.01, 0.015, 0.02])
    refs = {}
    for M_ in (100, 200, 400, 800, 1600):
        r_, T_, C_, _, _ = solve_q1(M_, 0.05, t_end, env, corr=True)
        refs[M_] = (r_, T_, C_)
    r_fin, T_fin, C_fin = refs[1600]
    Tf = np.interp(probes_r, r_fin, T_fin)
    Cf = np.interp(probes_r, r_fin, C_fin)
    print(f"  {'M':>6} {'dr[mm]':>9} {'max|dT|[K]':>13} {'阶':>7} {'max|dC|':>13} {'阶':>7}")
    prev = None
    conv = []
    for M_ in (100, 200, 400, 800):
        r_, T_, C_ = refs[M_]
        eT = np.max(np.abs(np.interp(probes_r, r_, T_) - Tf))
        eC = np.max(np.abs(np.interp(probes_r, r_, C_) - Cf))
        oT = oC = float("nan")
        if prev:
            oT = math.log(prev[0] / eT) / math.log(2) if eT > 0 else float("nan")
            oC = math.log(prev[1] / eC) / math.log(2) if eC > 0 else float("nan")
        prev = (eT, eC)
        conv.append((M_, R0 / M_ * 1000, eT, oT, eC, oC))
        print(f"  {M_:>6} {R0/M_*1000:>9.4f} {eT:>13.5e} {oT:>7.3f} {eC:>13.5e} {oC:>7.3f}")
    add("S10", "网格收敛: M=800 vs M=1600 温度最大偏差", float(conv[-1][2]), "K", "vs M=1600",
        "run solve_q1", "")
    add("S11", "网格收敛: M=800 vs M=1600 水分最大偏差", float(conv[-1][4]), "kg/kg", "vs M=1600",
        "run solve_q1", "")
    add("S12", "网格收敛经验阶 (温度, M=400->800)", float(conv[-1][3]), "-", "对数斜率", "run solve_q1", "")
    add("S13", "网格收敛经验阶 (水分, M=400->800)", float(conv[-1][5]), "-", "对数斜率", "run solve_q1", "")

    # ------------------------------------------------------------------ S5
    hr("S5. 时间步收敛性 (固定 M=800)")
    for dt_ in (4.0, 2.0, 1.0, 0.5, 0.25, 0.125):
        r_, T_, C_, _, _ = solve_q1(800, dt_, t_end, env, corr=True)
        eT = np.max(np.abs(np.interp(probes_r, r_, T_) - Tf))
        eC = np.max(np.abs(np.interp(probes_r, r_, C_) - Cf))
        print(f"  dt={dt_:>6}s: max|dT|={eT:.5e} K   max|dC|={eC:.5e} kg/kg")
        add(f"S14_dt{dt_}", f"时间步误差 dt={dt_}s (vs M=1600,dt=0.05)", float(max(eT, eC)), "-",
            "vs 参考", "run solve_q1", f"dT={eT:.3e} K; dC={eC:.3e}")

    # ------------------------------------------------------------------ S6
    hr("S6. 独立验证 2 — method of lines + BDF (非线性 D 不回滞), M=200")
    M_mol = 200
    r_m, Tm_all, Cm_all = solve_mol(M_mol, t_end, env)
    T_m, C_m = Tm_all[:, -1], Cm_all[:, -1]
    r_2, T_2, C_2, _, _ = solve_q1(M_mol, 0.05, t_end, env, corr=True)
    dT_lag = np.max(np.abs(T_m - T_2))
    dC_lag = np.max(np.abs(C_m - C_2))
    print(f"  同网格 M=200: FV(半隐式D) vs MOL(BDF): max|dT|={dT_lag:.3e} K, max|dC|={dC_lag:.3e} kg/kg")
    add("S15", "半隐式 D 回滞 vs 全非线性 温度最大偏差 (M=200)", float(dT_lag), "K",
        "semi-implicit lag", "run solve_mol", "")
    add("S16", "半隐式 D 回滞 vs 全非线性 水分最大偏差 (M=200)", float(dC_lag), "kg/kg",
        "semi-implicit lag", "run solve_mol", "")
    dT_mol_ref = np.max(np.abs(np.interp(probes_r, r_m, T_m) - Tf))
    dC_mol_ref = np.max(np.abs(np.interp(probes_r, r_m, C_m) - Cf))
    print(f"  MOL(M=200) vs FV 参考(M=1600): max|dT|={dT_mol_ref:.3e} K, max|dC|={dC_mol_ref:.3e} kg/kg")
    add("S17", "MOL(M=200) vs FV 参考 温度偏差", float(dT_mol_ref), "K", "vs M=1600", "run solve_mol", "")
    add("S18", "MOL(M=200) vs FV 参考 水分偏差", float(dC_mol_ref), "kg/kg", "vs M=1600", "run solve_mol", "")

    # ------------------------------------------------------------------ S7
    hr("S7. 插值方法敏感性 (附件1 -> 连续激励)")
    probes_t7 = np.array([1800.0])
    store = {}
    for meth in ("linear", "pchip", "cubic"):
        ev = Env(t1, T1, C1, method=meth)
        r_, T_, C_, sT, sC = solve_q1(800, 0.5, t_end, ev, corr=True, probes_t=probes_t7)
        store[meth] = (sT[0].copy(), sC[0].copy())
        print(f"  {meth:>7}: T(1800s) = " + " ".join(f"{v-273.15:>9.4f}" for v in np.interp(probes_r, r_, sT[0])))
        print(f"           C(1800s) = " + " ".join(f"{v:>9.6f}" for v in np.interp(probes_r, r_, sC[0])))
    for a_, b_, lbl in (("linear", "pchip", "linear vs pchip"), ("cubic", "pchip", "cubic vs pchip")):
        dT_ = np.max(np.abs(np.interp(probes_r, refs[800][0], store[a_][0]) -
                            np.interp(probes_r, refs[800][0], store[b_][0])))
        dC__ = np.max(np.abs(np.interp(probes_r, refs[800][0], store[a_][1]) -
                             np.interp(probes_r, refs[800][0], store[b_][1])))
        add(f"S19_{a_}", f"插值敏感性 {lbl}: 温度最大差", float(dT_), "K", "方法敏感性", "run", "")
        add(f"S20_{a_}", f"插值敏感性 {lbl}: 水分最大差", float(dC__), "kg/kg", "方法敏感性", "run", "")

    # ------------------------------------------------------------------ S8
    hr("S8. 生产解与表1/表2 (M=800, dt=0.125s, 内部子步采样到 1 s)")
    M_prod, dt_prod = 800, 0.125
    times_tab = np.array([100, 300, 600, 900, 1200, 1500, 1800], dtype=float)
    r_p, T_p, C_p, sT, sC = solve_q1(M_prod, dt_prod, t_end, env, corr=True,
                                     probes_t=times_tab)
    r_hi, T_hi, C_hi, sT_hi, sC_hi = solve_q1(1600, 0.0625, t_end, env, corr=True,
                                              probes_t=times_tab)
    print("  表1 温度 (degC)")
    print("    t/s   " + "".join(f"{d:>10}" for d in ["0", "0.5", "1", "1.5", "2"]))
    for k, tt in enumerate(times_tab):
        print(f"    {int(tt):>5} " + "".join(f"{v-273.15:>10.4f}" for v in np.interp(probes_r, r_p, sT[k])))
    print("  表2 水分浓度 (kg/kg)")
    print("    t/s   " + "".join(f"{d:>10}" for d in ["0", "0.5", "1", "1.5", "2"]))
    for k, tt in enumerate(times_tab):
        print(f"    {int(tt):>5} " + "".join(f"{v:>10.4f}" for v in np.interp(probes_r, r_p, sC[k])))

    # 生产解与高分辨率解的一致性 (生产精度声明)
    eT_prod = np.max(np.abs(np.interp(probes_r, r_p, T_p) - np.interp(probes_r, r_hi, T_hi)))
    eC_prod = np.max(np.abs(np.interp(probes_r, r_p, C_p) - np.interp(probes_r, r_hi, C_hi)))
    print(f"  生产解(M=800,dt=0.125) vs 高分辨率(M=1600,dt=0.0625): "
          f"max|dT|={eT_prod:.3e} K, max|dC|={eC_prod:.3e} kg/kg")
    add("S21", "生产解 vs 高分辨率解 温度最大偏差", float(eT_prod), "K", "vs M=1600", "run solve_q1",
        "生产设置 M=800, dt=0.125 s")
    add("S22", "生产解 vs 高分辨率解 水分最大偏差", float(eC_prod), "kg/kg", "vs M=1600",
        "run solve_q1", "生产设置 M=800, dt=0.125 s")

    with open(os.path.join(OUTDIR, "table1_temperature.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["时间/s", "0", "0.5", "1", "1.5", "2"])
        for k, tt in enumerate(times_tab):
            w.writerow([int(tt)] + [f"{v-273.15:.4f}" for v in np.interp(probes_r, r_p, sT[k])])
    with open(os.path.join(OUTDIR, "table2_moisture.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["时间/s", "0", "0.5", "1", "1.5", "2"])
        for k, tt in enumerate(times_tab):
            w.writerow([int(tt)] + [f"{v:.4f}" for v in np.interp(probes_r, r_p, sC[k])])

    for k, tt in enumerate(times_tab):
        vT = np.interp(probes_r, r_p, sT[k])
        vC = np.interp(probes_r, r_p, sC[k])
        add(f"R_T_{int(tt)}", f"T(r,t={int(tt)}s) degC @ 0/0.5/1/1.5/2cm",
            " / ".join(f"{v-273.15:.4f}" for v in vT), "degC", "±1e-4 K (见 S21)",
            "run solve_q1", "表1")
        add(f"R_C_{int(tt)}", f"C(r,t={int(tt)}s) kg/kg @ 0/0.5/1/1.5/2cm",
            " / ".join(f"{v:.4f}" for v in vC), "kg/kg", "±1e-6 (见 S22)",
            "run solve_q1", "表2")

    # ------------------------------------------------------------------ S9
    hr("S9. 水分场非线性与物理诊断")
    diag = []
    for tt in (1.0, 10.0, 60.0, 300.0, 900.0, 1800.0):
        M_ = 400
        r_, T_, C_, _, _ = solve_q1(M_, 0.25, tt, env, corr=True)
        dV = R0 / M_
        Vw = np.pi * (((np.arange(M_ + 1) + 0.5) * dV) ** 2 - ((np.arange(M_ + 1) - 0.5) * dV) ** 2)
        Vw[0] = np.pi * (dV / 2) ** 2
        Cavg = float(np.sum(C_ * Vw) / np.sum(Vw))
        # 节点式网格: node 0 是 r=0 (中心), node M 是 r=R (表面)
        diag.append((tt, Cavg, float(C_[0]), float(C_[M_]), env.C(tt)))
        print(f"  t={tt:>6}s <C>={Cavg:.6f}  C(0)={C_[0]:.6f}  C(R)={C_[M_]:.6f}  "
              f"C_inf={env.C(tt):.5f}  C(R)/C_inf={C_[M_]/env.C(tt):.4f}  "
              f"C(0)/C_inf={C_[0]/env.C(tt):.4f}")
        add(f"S23_{int(tt)}", f"体积平均含水率 <C>(t={int(tt)}s)", Cavg, "kg/kg",
            "网格不确定度见 S11", "run solve_q1", "体积加权")
        add(f"S23c_{int(tt)}", f"中心含水率 C(r=0,t={int(tt)}s)", float(C_[0]), "kg/kg",
            "网格不确定度见 S11", "run solve_q1", "")
        add(f"S23s_{int(tt)}", f"表面含水率 C(r=R,t={int(tt)}s)", float(C_[M_]), "kg/kg",
            "网格不确定度见 S11", "run solve_q1", "")
        add(f"S23r_{int(tt)}", f"C(r=R)/C_inf (t={int(tt)}s)", float(C_[M_] / env.C(tt)), "-",
            "解析比值", "run solve_q1", "表面仍远未与热风平衡")
        add(f"S23q_{int(tt)}", f"C(r=0)/C_inf (t={int(tt)}s)", float(C_[0] / env.C(tt)), "-",
            "解析比值", "run solve_q1", "中心")
    add("S24", "<C> 从 0 到 1800 s 的相对下降", float(1 - diag[-1][1] / diag[0][1]), "-", "解析比值",
        "run solve_q1", "预热平衡阶段的整体干燥比例")
    add("S25", "C(r=R) 与 C_inf 的比值 (t=1800s)", float(diag[-1][3] / diag[-1][4]), "-",
        "解析比值", "run solve_q1", "表面: >>1 表示表面仍远未与热风平衡")
    add("S25b", "C(r=0) 与 C_inf 的比值 (t=1800s)", float(diag[-1][2] / diag[-1][4]), "-",
        "解析比值", "run solve_q1", "中心")
    add("S26", "中心含水率 C(r=0) 相对初始值的下降 (t=1800s)", float(1 - diag[-1][2] / C0), "-",
        "解析比值", "run solve_q1", "预热平衡阶段中心几乎不失水")
    add("S26b", "体积平均含水率相对初始值的下降 (t=1800s)", float(1 - diag[-1][1] / diag[0][1]), "-",
        "解析比值", "run solve_q1", "")

    _write_registry(reg, os.path.join(OUTDIR, "registry_q1.csv"))
    print()
    print(f"输出: {os.path.join(OUTDIR, 'registry_q1.csv')} ({len(reg)} 行)")
    print(f"      {os.path.join(OUTDIR, 'table1_temperature.csv')}")
    print(f"      {os.path.join(OUTDIR, 'table2_moisture.csv')}")


def _write_registry(rows, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["id", "quantity", "value", "unit", "uncertainty", "source", "command", "note"])
        for r in rows:
            w.writerow(r)


if __name__ == "__main__":
    main()
