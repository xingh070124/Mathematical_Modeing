# -*- coding: utf-8 -*-
"""Adversarial probes for the Q2 tex section:

P1. Is the maximum principle's lower bound exactly recovered when the
    evaporation term is switched off (paper: "精确回到 28.0000 degC, 越界 0.0000 K")?
    Run the real solver, two resolutions, full 10800 s.
P2. What is the run's own volume-mean dT/dt, versus the 1e-3 K/s assumed in
    registry_q2_energy EC_main to build the 0.2520 ratio?
P3. Per-column count of temperature non-monotone samples, to test the paper's
    "温度的非单调点全部集中在 r>=1.5 cm" claim.
"""
from __future__ import annotations

import os
import sys

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl  # noqa: E402
from q1_solve import Env, load_attachment1  # noqa: E402
from q2_solve import Par, march, out_indices, T0K  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "outputs")

# --- P3 + P2: from the delivered result2.xlsx -------------------------------
wb = openpyxl.load_workbook(os.path.join(OUT, "result2.xlsx"),
                            data_only=True, read_only=True)
rh = np.array([float(v) for v in next(wb["温度"].iter_rows(values_only=True))[1:]])
rowsT = list(wb["温度"].iter_rows(values_only=True))[1:]
rowsC = list(wb["水分浓度"].iter_rows(values_only=True))[1:]
wb.close()
T = np.array([[float(x) for x in r[1:]] for r in rowsT])
C = np.array([[float(x) for x in r[1:]] for r in rowsC])
print("=" * 100)
print("P3  温度逐列时间下降点数 (阈值 1e-4, 与 q2_derive.py 同判据)")
print("=" * 100)
nT = [int(np.sum(np.diff(T[:, j]) < -1e-4)) for j in range(T.shape[1])]
for j, n in enumerate(nT):
    print(f"  r={rh[j]:4.1f} cm : {n:5d}")
print(f"  合计 = {sum(nT)};  r>=1.5 cm 合计 = {sum(nT[15:])};  "
      f"r<=1.4 cm 合计 = {sum(nT[:15])}")
print(f"  论文称'温度的非单调点全部集中在 r>=1.5 cm 的表层' -> "
      f"r in [1.1,1.4] cm 实有 {sum(nT[11:15])} 点")

print()
print("=" * 100)
print("P2  运行解自带的体积平均 dT/dt (梯形权, 与 FG_Cbar 同一求积)")
print("=" * 100)
w = np.full(21, 1.0)
w[0] = w[-1] = 0.5
w *= (0.02 / 20.0) * (rh * 1e-2)
Tbar = T @ w / w.sum()
dt = np.gradient(Tbar)          # per second
for t0 in (1800, 3600, 5400, 7200, 9000, 10800):
    print(f"  t={t0:6d} s: dTbar/dt = {dt[t0-1]*1e3:8.4f} e-3 K/s")
print(f"  时均 (Tbar(10800)-Tbar(0))/10800 = "
      f"{(Tbar[-1]-Tbar[0])/10800*1e3:.4f} e-3 K/s")
print(f"  最大 |dTbar/dt| = {np.abs(dt).max()*1e3:.3f} e-3 K/s @ "
      f"t={int(np.argmax(np.abs(dt)))+1} s")
print(f"  1~2 h 平均 = {(Tbar[7200-1]-Tbar[3600-1])/3600*1e3:.4f} e-3 K/s")
print(f"  EC_main 假设的 dT/dt = 1.0 e-3 K/s  ->  "
      f"偏小 {((Tbar[-1]-Tbar[0])/10800)/1e-3:.2f} 倍 (时均) / "
      f"{np.abs(dt).max()/1e-3:.2f} 倍 (峰值)")

print()
print("=" * 100)
print("P1  关闭蒸发项后温度下界是否精确等于 28 degC (真实求解器)")
print("=" * 100)
t1, TT, CC = load_attachment1()
env = Env(t1, TT, CC, method="pchip")
for M, dtt in ((200, 0.25), (400, 1.0 / 16)):
    par = Par(hevap="off")
    o = march(M, dtt, 10800.0, env, par, out_idx=out_indices(M))
    Ts = o["T_snap"]
    mn = float(Ts.min()) - 273.15
    print(f"  H_evap=0, M={M}, dt={dtt:.6f}: min T = {mn:.10f} degC, "
          f"越界 = {28.0 - mn:.3e} K, 最大值 = {Ts.max()-273.15:.4f} degC")
    par2 = Par()
    o2 = march(M, dtt, 10800.0, env, par2, out_idx=out_indices(M))
    Ts2 = o2["T_snap"]
    print(f"  含蒸发 , M={M}, dt={dtt:.6f}: min T = {Ts2.min()-273.15:.10f} degC, "
          f"越界 = {28.0-(Ts2.min()-273.15):.6f} K, "
          f"max = {Ts2.max()-273.15:.6f} degC")
    print(f"  含/不含蒸发 max|dT| = {np.max(np.abs(Ts2-Ts)):.6f} K")
