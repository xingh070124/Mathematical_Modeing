# -*- coding: utf-8 -*-
"""Audit H: the 7.4 monotonicity table: threshold sensitivity (doc uses diff < -1e-4)."""
from __future__ import annotations
import os, sys
import numpy as np
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
wb = load_workbook(os.path.join(ROOT, "outputs", "result2.xlsx"), read_only=True, data_only=True)
rows = list(wb["温度"].iter_rows(min_row=2, values_only=True))
T = np.array([[float(x) for x in r[1:]] for r in rows])
rad = np.array([float(x) for x in next(wb["温度"].iter_rows(min_row=1, max_row=1, values_only=True))[1:]])
wb.close()
print("doc §7.4 uses threshold diff < -1e-4 in q2_derive.py line 254")
for thr in (0.0, -1e-7, -1e-6, -1e-5, -1e-4, -1e-3):
    n = [(np.diff(T[:, j]) < thr).sum() for j in range(T.shape[1])]
    print(f"  thr={thr:>8}: j15(r=1.5)={n[15]:5d}  j19(r=1.9)={n[19]:5d}  j20(r=2.0)={n[20]:5d}"
          f"  outer={sum(n[15:]):5d}  inner(r<=1.0)={sum(n[:11]):4d}  total={sum(n):6d}")
print("\n  doc table (r=1.5,1.9,2.0) = 72, 632, 800 ; outer 2417 ; inner 0")
print("  my recomputation from the published result2.xlsx with thr=0 gives")
n0 = [(np.diff(T[:, j]) < 0).sum() for j in range(T.shape[1])]
print("   ", n0)
print(f"    -> r=1.5: {n0[15]}, r=1.9: {n0[19]}, r=2.0: {n0[20]}, outer {sum(n0[15:])}, "
      f"inner {sum(n0[:11])}, total {sum(n0)}")
