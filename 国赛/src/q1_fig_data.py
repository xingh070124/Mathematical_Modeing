"""
图2 环境激励数据可视化; 图3 问题1 温度场/水分场结果可视化.

figure contract:
  fig_q1_env      核心结论: 附件1 给出温升/湿增的时变激励, C_inf 全程含测量噪声
                  (240 步中 80 步下降), 但问题1 所用的 0~1800 s 窗口内单调 --
                  这决定了插值方法的选择 (PCHIP 保形) 与问题1 对插值不敏感.
  fig_q1_fields   核心结论: 30 min 内温度全场快速响应 (径向梯度明显), 水分仅在
                  表层脱湿 -- 热/湿时间尺度分离 (R^2/alpha : R^2/D(C0) ~ 1:34)
                  的直接可视化证据.

数据源: A题/附件/附件1.xlsx (环境激励), outputs/result1.xlsx (生产解 M=3200, dt=2^-8 s).

运行:  python src/q1_fig_data.py
输出:  paper/figures/fig_q1_env.pdf|.png, paper/figures/fig_q1_fields.pdf|.png
"""

from __future__ import annotations

import os
import sys

import matplotlib.pyplot as plt
import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fig_style import C as PALETTE, GREY, SERIES, save, setup  # noqa: E402
from q1_sensitivity import robust_sigma  # noqa: E402
from q1_solve import load_attachment1  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "outputs", "result1.xlsx")
T_TAB = [100, 300, 600, 900, 1200, 1500, 1800]


def load_result1():
    from openpyxl import load_workbook

    wb = load_workbook(XLSX, data_only=True, read_only=True)

    def grid(ws):
        rows = list(ws.iter_rows(values_only=True))
        hdr = rows[0]
        t = np.array([r[0] for r in rows[1:]], dtype=float)
        v = np.array([[float(x) for x in r[1:]] for r in rows[1:]], dtype=float)
        return np.array([float(h) for h in hdr[1:]]), t, v

    rh, tt, T = grid(wb["温度"])
    _, tc, C = grid(wb["水分浓度"])
    wb.close()
    assert np.array_equal(tt, tc)
    return rh, tt, T, C


def fig_env():
    t1, T1, C1 = load_attachment1()
    sigT, sigC = robust_sigma(T1), robust_sigma(C1)
    n_drop = int(np.sum(np.diff(C1) < 0))

    fig, axes = plt.subplots(1, 2, figsize=(7.1, 2.7))

    # ---- (a) 烘房温度 ----
    ax = axes[0]
    ax.axvspan(0, 1.8, color=PALETTE["band"], zorder=0)
    ax.text(7.2, 52.5, "问题1 时段 (预热平衡, 0–1800 s)", ha="center",
            va="top", fontsize=7, color=PALETTE["orange"])
    ax.plot(t1 / 1e3, T1, "o", ms=2.2, mfc=PALETTE["blue"], mec="none", zorder=2,
            label="附件1 采样 (Δt=60 s)")
    tt = np.linspace(0, t1[-1], 1200)
    from scipy.interpolate import PchipInterpolator
    ax.plot(tt / 1e3, PchipInterpolator(t1, T1)(tt), "-", color=GREY["ink"],
            lw=1.0, zorder=3, label="PCHIP 连续化")
    ax.annotate("色带：问题1 时段 ($0$–$1.8\\times10^3$ s)\n$T_\\infty$: 28.000 → 50.246 $^\\circ$C",
                xy=(5.0, 38.0), fontsize=7.5)
    ax.set_xlabel("时间 $t$ (×10$^3$ s)")
    ax.set_ylabel("烘房温度 $T_\\infty$ (°C)")
    ax.set_ylim(26, 54)
    ax.set_xlim(-0.4, 14.8)
    ax.legend(loc="lower right")
    ax.set_title("(a) 烘房温度（附件1）", fontsize=8.5)

    # ---- (b) 烘房水分浓度 ----
    ax = axes[1]
    ax.axvspan(0, 1.8, color=PALETTE["band"], zorder=0)
    ax.plot(t1 / 1e3, C1, "o", ms=2.2, mfc=PALETTE["orange"], mec="none", zorder=2)
    drop_idx = np.where(np.diff(C1) < 0)[0] + 1
    ax.plot(t1[drop_idx] / 1e3, C1[drop_idx], "o", ms=4.5, mfc="none",
            mec=PALETTE["red"], mew=0.8, zorder=3, label=f"下降步 ({n_drop}/240 步, 噪声)")
    tt = np.linspace(0, t1[-1], 1200)
    ax.plot(tt / 1e3, PchipInterpolator(t1, C1)(tt), "-", color=GREY["ink"],
            lw=1.0, zorder=4, label="PCHIP 连续化")
    ax.annotate("色带：问题1 时段（区间内单调递增）",
                xy=(4.6, 0.0545), fontsize=7.3, va="top", color=PALETTE["orange"])
    ax.annotate("$C_\\infty$: 0.0196 → 0.0503 kg/kg\n"
                f"噪声 $\\sigma_C\\approx{sigC * 1e4:.1f}\\times10^{{-4}}$ kg/kg"
                f"（$\\sigma_T\\approx{sigT * 1e3:.0f}\\times10^{{-3}}$ °C）",
                xy=(5.2, 0.0215), fontsize=7.5)
    ax.set_xlabel("时间 $t$ (×10$^3$ s)")
    ax.set_ylabel("烘房水分浓度 $C_\\infty$ (kg/kg)")
    ax.set_ylim(0.014, 0.057)
    ax.set_xlim(-0.4, 14.8)
    ax.legend(loc="lower right")
    ax.set_title("(b) 烘房水分浓度（附件1，含测量噪声）", fontsize=8.5)

    fig.tight_layout()
    save(fig, "fig_q1_env")


def fig_fields(rh, tt, T, C):
    cmap_T = plt.get_cmap("inferno")       # 温度: 热
    cmap_C = plt.get_cmap("viridis")       # 水分: 冷
    t_idx = [int(np.where(tt == v)[0][0]) for v in T_TAB]
    shades = [plt.cm.viridis(0.02 + 0.78 * k / (len(T_TAB) - 1))
              for k in range(len(T_TAB))]

    fig, axes = plt.subplots(2, 2, figsize=(7.1, 5.4))

    # ---- (a) 温度场热图 ----
    ax = axes[0, 0]
    Tm = ax.pcolormesh(tt, rh, T.T, cmap=cmap_T, shading="auto",
                       vmin=T.min(), vmax=T.max())
    cb = fig.colorbar(Tm, ax=ax, pad=0.015)
    cb.set_label("温度 $T$ (°C)", fontsize=7.5)
    cb.ax.tick_params(labelsize=7)
    ax.set_xlabel("时间 $t$ (s)")
    ax.set_ylabel("到中心距离 $r$ (cm)")
    ax.set_title("(a) 温度场 $T(r,t)$", fontsize=8.5)

    # ---- (b) 水分场热图 ----
    ax = axes[0, 1]
    Cm = ax.pcolormesh(tt, rh, C.T, cmap=cmap_C, shading="auto",
                       vmin=C.min(), vmax=C.max())
    cb = fig.colorbar(Cm, ax=ax, pad=0.015)
    cb.set_label("水分浓度 $C$ (kg/kg)", fontsize=7.5)
    cb.ax.tick_params(labelsize=7)
    ax.set_xlabel("时间 $t$ (s)")
    ax.set_ylabel("到中心距离 $r$ (cm)")
    ax.set_title("(b) 水分浓度场 $C(r,t)$", fontsize=8.5)

    # ---- (c) 温度径向剖面 ----
    ax = axes[1, 0]
    for k, i in enumerate(t_idx):
        ax.plot(rh, T[i, :], "-", color=shades[k], lw=1.2)
        ax.text(2.06, T[i, -1], f"{T_TAB[k]}", fontsize=6.8,
                va="center", color=shades[k])
    ax.text(0.08, 36.3, "线色由深至浅：$t=100\\to1800$ s", fontsize=7,
            color=GREY["dark"])
    ax.set_xlim(0, 2.75)
    ax.set_xlabel("到中心距离 $r$ (cm)")
    ax.set_ylabel("温度 $T$ (°C)")
    ax.set_title("(c) 温度径向剖面（表1 时刻）", fontsize=8.5)

    # ---- (d) 水分径向剖面 ----
    ax = axes[1, 1]
    for k, i in enumerate(t_idx):
        ax.plot(rh, C[i, :], "-", color=shades[k], lw=1.2)
        ax.text(2.06, C[i, -1], f"{T_TAB[k]}", fontsize=6.8,
                va="center", color=shades[k])
    ax.text(2.42, 1.68, "t/s →", fontsize=7, color=GREY["dark"])
    ax.axhline(2.55, color=PALETTE["red"], lw=0.7, ls=(0, (4, 2)))
    ax.text(0.06, 2.575, "$C_0=2.55$", fontsize=7, color=PALETTE["red"])
    ax.set_xlim(0, 2.75)
    ax.set_xlabel("到中心距离 $r$ (cm)")
    ax.set_ylabel("水分浓度 $C$ (kg/kg)")
    ax.set_title("(d) 水分浓度径向剖面（表2 时刻）", fontsize=8.5)

    fig.tight_layout()
    save(fig, "fig_q1_fields")


def main():
    setup()
    fig_env()
    rh, tt, T, C = load_result1()
    fig_fields(rh, tt, T, C)


if __name__ == "__main__":
    main()
