"""
AUDIT: isolate the discrepancy between
  (a) the report's semi-analytic reference (src/feishu_compare.py series_T), and
  (b) my independently built one (src/audit_feishu_math.py).

Nothing is written by feishu_compare.py -- its functions are imported and called
in-process; main() is never executed, so outputs/ is not touched.

Run: python src/audit_feishu_series.py
Out: outputs/audit_feishu_series.log
"""

from __future__ import annotations

import math
import os
import sys

import numpy as np
from scipy.interpolate import PchipInterpolator
from scipy.optimize import brentq
from scipy.special import j0 as sj0
from scipy.special import j1 as sj1

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))
import feishu_compare as FC  # noqa: E402  (import only; main() not called)

ROWS = []


def say(s=""):
    ROWS.append(s)
    print(s, flush=True)


ALPHA = FC.ALPHA
R0 = FC.R0
T0_K = FC.T0_K

say("=" * 78)
say("S1. eigen roots: theirs (FC.eigen) vs mine (brentq scan)")
say("=" * 78)
xt, lamt, ct = FC.eigen(FC.BI, 400)
say("theirs: n=%d  x[:6]=" % len(xt) + ", ".join("%.10f" % v for v in xt[:6]))
say("        x[-1]=%.10f" % xt[-1])


def mine(nmodes, bi=FC.BI, step=0.02):
    f = lambda x: x * sj1(x) - bi * sj0(x)
    xs = []
    a = 1e-9
    fa = f(a)
    x = a
    while len(xs) < nmodes and x < 4 * nmodes + 60:
        x2 = x + step
        fb = f(x2)
        if fa * fb < 0:
            xs.append(brentq(f, x, x2, xtol=1e-15, rtol=8.9e-16))
        x, fa = x2, fb
    return np.array(xs)


xm = mine(400)
say("mine  : n=%d  x[:6]=" % len(xm) + ", ".join("%.10f" % v for v in xm[:6]))
say("        x[-1]=%.10f" % xm[-1])
n = min(len(xt), len(xm))
say("max|theirs - mine| over %d common roots = %.3e" % (n, np.max(np.abs(xt[:n] - xm[:n]))))
say("")
say("residual of THEIR roots under x J1 - Bi J0:  max=%.3e"
    % np.max(np.abs(xt * sj1(xt) - FC.BI * sj0(xt))))
say("residual of MY roots:                       max=%.3e"
    % np.max(np.abs(xm * sj1(xm) - FC.BI * sj0(xm))))
say("")
say("check the LARGE-x tail identically (roots must be ~pi apart):")
d_t = np.diff(xt[-20:])
d_m = np.diff(xm[-20:])
say("  theirs last-20 spacing min/max = %.6f / %.6f  (pi = %.6f)"
    % (d_t.min(), d_t.max(), math.pi))
say("  mine   last-20 spacing min/max = %.6f / %.6f" % (d_m.min(), d_m.max()))

say("")
say("=" * 78)
say("S2. expansion coefficients")
say("=" * 78)
say("theirs c[:6] = " + ", ".join("%.10f" % v for v in ct[:6]))
cm = 2 * sj0(xm) / (xm * (sj0(xm) ** 2 + sj1(xm) ** 2))
say("mine   c[:6] = " + ", ".join("%.10f" % v for v in cm[:6]))
say("max|ct - cm| = %.3e" % np.max(np.abs(ct[:n] - cm[:n])))
say("")
say("their formula line 117:  c = 2*j1(x)/(x*(j0(x)**2 + j1(x)**2))   <- read it")
import inspect  # noqa: E402
src = inspect.getsource(FC.eigen)
say("--- FC.eigen source ---")
say(src)
say("--- FC.series_T source ---")
say(inspect.getsource(FC.series_T))

say("")
say("=" * 78)
say("S3. the decisive diagnostic: how well does sum_n c_n J0(lam_n r) == 1 hold?")
say("=" * 78)
for N in (8, 20, 50, 100, 200, 400):
    s0 = float(np.sum(ct[:N] * 1.0))
    say("  N=%3d  sum c_n ( = sum c_n J0 at r=0 ) = %.10f   1-sum = %+.3e"
        % (N, s0, 1.0 - s0))
say("")
say("  => this is the TRUNCATION of the forcing representation.  A deficit here")
say("     multiplies Tinf'(t) in  b_n' + mu_n b_n = -c_n Tinf',  so it biases the")
say("     whole solution by roughly (1-sum c_n) x (Tinf(1800)-Tinf(0)) x response.")

say("")
say("=" * 78)
say("S4. real case: their series_T vs my series_at vs the numeric CN solution")
say("=" * 78)
from openpyxl import load_workbook  # noqa: E402
ATT1 = os.path.join(ROOT, "A题", "附件", "附件1.xlsx")
wb = load_workbook(ATT1, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
rr = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
wb.close()
arr = np.array([[float(v) for v in r[:3]] for r in rr])
t_env, T_env_C = arr[:, 0], arr[:, 1]
Tf = PchipInterpolator(t_env, T_env_C)
t_out = np.arange(0.0, 1800.0 + 1e-9, 1.0)
Tinf_out = np.asarray(Tf(t_out) + 273.15, float)
say("Tinf(0)=%.6f K  Tinf(1800)=%.6f K  (so T0-Tinf(0)=%.3e)"
    % (Tinf_out[0], Tinf_out[-1], T0_K - Tinf_out[0]))

Mnum, dtnum = 1600, 2.0 ** -5
r = np.arange(Mnum + 1) * (R0 / Mnum)
Tnum = FC.solve_T(Mnum, dtnum, t_out, lambda t: float(Tf(t)) + 273.15,
                  0.5, "exact_cn")[1]
say("")
say("report's numeric  T(1800s) centre/surface = %.10f / %.10f" % (Tnum[-1, 0], Tnum[-1, -1]))
say("registry F43/F44                            = 306.72583087 / 309.93578369")

Ta_num, _, _ = FC.series_T(t_out, Tinf_out, r, n_modes=400)
say("")
say("report's semi-analytic series_T (400 modes):")
say("  T(1800s) centre/surface = %.10f / %.10f" % (Ta_num[-1, 0], Ta_num[-1, -1]))
say("  max|Tnum - Tana| = %.10e K" % np.abs(Tnum - Ta_num).max())


def series_mine(N, t_out, Tinf_out, r_eval):
    b = ct[:N] * (T0_K - Tinf_out[0])
    mu = ALPHA * (xt[:N] / R0) ** 2
    cc = ct[:N]
    out = np.empty((len(t_out), np.size(r_eval)))
    J = sj0(np.outer(xt[:N] / R0, np.atleast_1d(np.asarray(r_eval, float))))
    out[0] = Tinf_out[0] + b @ J
    for k in range(len(t_out) - 1):
        d = t_out[k + 1] - t_out[k]
        s = (Tinf_out[k + 1] - Tinf_out[k]) / d
        E = np.exp(-mu * d)
        b = E * b - cc * s * (1 - E) / mu
        out[k + 1] = Tinf_out[k + 1] + b @ J
    return out


Ta_mine = series_mine(400, t_out, Tinf_out, r)
say("")
say("my own re-implementation, SAME 400 roots, SAME coefficients, SAME forcing:")
say("  T(1800s) centre/surface = %.10f / %.10f" % (Ta_mine[-1, 0], Ta_mine[-1, -1]))
say("  max|Tnum - Ta_mine| = %.10e K" % np.abs(Tnum - Ta_mine).max())
say("  max|Ta_num - Ta_mine| = %.3e K   <- should be ~0 if the two are the same object"
    % np.abs(Ta_num - Ta_mine).max())

say("")
say("=" * 78)
say("S5. mode-count convergence of the report's own series_T")
say("=" * 78)
say("  N_modes   T(1800,r=0)        T(1800,r=R)        max|Tnum - Tana|")
for N in (8, 20, 50, 100, 200, 400, 800):
    try:
        Tn, _, _ = FC.series_T(t_out, Tinf_out, r, n_modes=N)
        say("  %5d     %.10f    %.10f   %.6e"
            % (N, Tn[-1, 0], Tn[-1, -1], np.abs(Tnum - Tn).max()))
    except Exception as e:  # noqa: BLE001
        say("  %5d     FAILED: %s" % (N, e))
say("")
say("If the mode count is the controlling error, the agreement at N=400 is")
say("coincidental and N=800 must move it by the same order.")

say("")
say("=" * 78)
say("S6. does the ALPHA inside series_T equal the ALPHA inside solve_T?")
say("=" * 78)
say("FC.ALPHA = %.17g   K/(RHO*CP) = %.17g" % (ALPHA, FC.K_COND / (FC.RHO * FC.CP)))
say("FC.T0_K  = %.6f" % T0_K)

os.makedirs(os.path.join(ROOT, "outputs"), exist_ok=True)
with open(os.path.join(ROOT, "outputs", "audit_feishu_series.log"), "w",
          encoding="utf-8") as f:
    f.write("\n".join(ROWS) + "\n")
print("\nwrote outputs/audit_feishu_series.log")
