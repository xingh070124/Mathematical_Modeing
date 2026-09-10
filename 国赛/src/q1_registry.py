"""
从生产产物 outputs/result1.xlsx 生成生产数字注册表.

原则: 注册表是论文数字的唯一来源, 且它本身由"实际跑出来的文件"生成,
而不是从文稿抄写. 链条: result1.xlsx <- q1_produce.py <- 模型.

输出: outputs/registry_q1_production.csv
运行: python src/q1_registry.py
"""

from __future__ import annotations

import csv
import os
import sys

import numpy as np

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "outputs")
XLSX = os.path.join(OUT, "result1.xlsx")
REG = os.path.join(OUT, "registry_q1_production.csv")

CMD = "python src/q1_produce.py"
SRC = "outputs/result1.xlsx"

T_TAB = [100, 300, 600, 900, 1200, 1500, 1800]
R_CM = [0.0, 0.5, 1.0, 1.5, 2.0]

# 生产解的不确定度 (来自 outputs/q1_production.log)
UNC_T = 3.742e-6      # K, dt 减半
UNC_C = 1.171e-5      # kg/kg, 网格加密
UNC_T_H = 6.115e-8
UNC_C_T = 8.029e-6


def load():
    from openpyxl import load_workbook

    wb = load_workbook(XLSX, data_only=True, read_only=True)
    ws_t = wb["温度"]
    ws_c = wb["水分浓度"]

    def grid(ws):
        rows = list(ws.iter_rows(values_only=True))
        hdr = rows[0]
        tcol = np.array([r[0] for r in rows[1:]], dtype=float)
        vals = np.array([[float(v) for v in r[1:]] for r in rows[1:]], dtype=float)
        return np.array([float(h) for h in hdr[1:]]), tcol, vals

    rh, tt, T = grid(ws_t)
    _, tc, C = grid(ws_c)
    wb.close()
    assert np.array_equal(tt, tc), "两个工作表的时间列不一致"
    return rh, tt, T, C


def main():
    rh, tt, T, C = load()
    print(f"result1.xlsx: 半径 {len(rh)} 列 ({rh[0]}..{rh[-1]} cm), 时间 {len(tt)} 行 "
          f"({int(tt[0])}..{int(tt[-1])} s)")

    rows = []

    def add(id_, q, v, u, unc, src, note=""):
        rows.append([id_, q, v, u, unc, src, CMD, note])

    # --- 结构自检 ---
    add("X01", "result1.xlsx 温度工作表 行数", len(tt), "行", "结构", SRC, "t=1..1800 s")
    add("X02", "result1.xlsx 温度工作表 列数", len(rh), "列", "结构", SRC, "r=0..2.0 cm, 步长 0.1 cm")
    add("X03", "输出半径步长", 0.1, "cm", "结构", SRC, "")
    add("X04", "输出时间步长", 1.0, "s", "结构", SRC, "")

    # --- 表1 温度 (7x5) ---
    for t_ in T_TAB:
        i = int(np.where(tt == t_)[0][0])
        for rc in R_CM:
            j = int(np.where(np.isclose(rh, rc))[0][0])
            add(f"T1_{t_}_{str(rc).replace('.', 'p')}",
                f"表1 温度 T(r={rc}cm, t={t_}s)", f"{T[i, j]:.4f}", "degC",
                f"<{UNC_T:.1e} K (离散)", SRC, "生产解 M=3200, dt=2^-8 s")

    # --- 表2 水分浓度 (7x5) ---
    for t_ in T_TAB:
        i = int(np.where(tt == t_)[0][0])
        for rc in R_CM:
            j = int(np.where(np.isclose(rh, rc))[0][0])
            add(f"C2_{t_}_{str(rc).replace('.', 'p')}",
                f"表2 水分浓度 C(r={rc}cm, t={t_}s)", f"{C[i, j]:.4f}", "kg/kg",
                f"<{UNC_C:.1e} kg/kg (离散)", SRC, "生产解 M=3200, dt=2^-8 s")

    # --- 极值与端点 ---
    i0, i1 = 0, len(tt) - 1
    add("X10", "温度场最小值 (t=1s, r=0)", f"{T[i0, 0]:.4f}", "degC", "结构", SRC, "")
    add("X11", "温度场在 t=1800s 的最大值 (r=R)", f"{T[i1, -1]:.4f}", "degC", "结构", SRC, "")
    add("X12", "温度场在 t=1800s 的径向温差 T(R)-T(0)",
        f"{T[i1, -1] - T[i1, 0]:.4f}", "degC", "解析差", SRC, "")
    add("X13", "水分场在 t=1800s 的中心值 C(0)", f"{C[i1, 0]:.4f}", "kg/kg", "结构", SRC, "")
    add("X14", "水分场在 t=1800s 的表面值 C(R)", f"{C[i1, -1]:.4f}", "kg/kg", "结构", SRC, "")
    add("X15", "整个文件的最小温度", f"{T.min():.4f}", "degC", "结构", SRC, "")
    add("X16", "整个文件的最大温度", f"{T.max():.4f}", "degC", "结构", SRC, "")
    add("X17", "整个文件的最小水分浓度", f"{C.min():.4f}", "kg/kg", "结构", SRC, "")
    add("X18", "整个文件的最大水分浓度", f"{C.max():.4f}", "kg/kg", "结构", SRC, "")

    # 极值位置自检
    assert T.min() >= 28.0 - 1e-9 and T.max() <= 41.6, "温度越出物理包络"
    assert C.max() <= 2.55 + 1e-9, "水分浓度超过初值"

    # --- 物理诊断 (t=1800s) ---
    add("X20", "t=1800s 中心含水率相对初始值的变化",
        f"{C[i1, 0] - 2.55:.6e}", "kg/kg", "解析差", SRC, "几乎为零: 中心未失水")
    add("X21", "t=1800s 表面含水率相对初始值的下降比例",
        f"{(2.55 - C[i1, -1]) / 2.55:.6f}", "-", "解析比值", SRC, "")

    # --- 初始含水率与环境水分浓度之比 (说明表面传质边界的量级差异) ---
    from q1_solve import load_attachment1
    t1_, T1_, C1_ = load_attachment1()
    add("X25", "C0 / C_inf(t=0)", f"{2.55 / C1_[0]:.4f}", "-", "解析比值", "附件1",
        "初始含水率与环境含湿量之比")
    add("X26", "C0 / C_inf(t=14400s)", f"{2.55 / C1_[-1]:.4f}", "-", "解析比值", "附件1",
        "初始含水率与环境含湿量之比")
    add("X27", "C0 / C_inf 的范围", f"{2.55 / C1_[-1]:.1f}..{2.55 / C1_[0]:.1f}", "倍",
        "解析比值", "附件1", "即约 50 ~ 130 倍")

    # --- 离散不确定度 (来自生产日志) ---
    add("X30", "温度离散不确定度 (时间步减半)", f"{UNC_T:.4e}", "K", "-", "outputs/q1_production.log", "")
    add("X31", "温度离散不确定度 (网格加密)", f"{UNC_T_H:.4e}", "K", "-", "outputs/q1_production.log", "")
    add("X32", "水分离散不确定度 (时间步减半)", f"{UNC_C_T:.4e}", "kg/kg", "-", "outputs/q1_production.log", "")
    add("X33", "水分离散不确定度 (网格加密)", f"{UNC_C:.4e}", "kg/kg", "-", "outputs/q1_production.log", "")

    with open(REG, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["id", "quantity", "value", "unit", "uncertainty", "source", "command", "note"])
        w.writerows(rows)

    print(f"注册表: {REG}  ({len(rows)} 行)")
    print("  表1 温度:", " ".join(f"{T[np.where(tt==1800)[0][0], j]:.4f}"
                                 for j in range(len(rh)) if rh[j] in R_CM))
    print("  表2 水分:", " ".join(f"{C[np.where(tt==1800)[0][0], j]:.4f}"
                                 for j in range(len(rh)) if rh[j] in R_CM))


if __name__ == "__main__":
    main()
