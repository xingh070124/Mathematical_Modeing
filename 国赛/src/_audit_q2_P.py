# -*- coding: utf-8 -*-
"""Audit P: re-run the PRODUCTION configuration (M=1600, dt=1/64) for t<=1800 s and
compare bit-for-bit against the shipped outputs/result2.xlsx rows 1..1800."""
from __future__ import annotations
import os, sys, time
import numpy as np
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
from q1_solve import Env, load_attachment1
from q2_solve import Par, march, out_indices

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
T_END = 1800.0
M, DT = 1600, 1.0 / 64.0

t1, T1, C1 = load_attachment1()
env = Env(t1, T1, C1, method="pchip")
idx = out_indices(M)
t0 = time.time()
o = march(M, DT, T_END, env, Par(), out_idx=idx)
print(f"[re-run] M={M}, dt={DT:.10g} s, t_end={T_END:.0f} s -> "
      f"{o['stats']['nsteps']} steps, {time.time()-t0:.1f} s", flush=True)

wb = load_workbook(os.path.join(ROOT, "outputs", "result2.xlsx"),
                   read_only=True, data_only=True)
def grid(sn):
    rows = list(wb[sn].iter_rows(min_row=2, values_only=True))
    return np.array([[float(x) for x in r[1:]] for r in rows])
Tx, Cx = grid("温度"), grid("水分浓度")
wb.close()

n = int(T_END)
Tr = o["T_snap"][:n] - 273.15
Cr = o["C_snap"][:n]
Tr_x, Cr_x = Tx[:n], Cx[:n]
dT = np.abs(Tr - Tr_x)
dC = np.abs(Cr - Cr_x)
print(f"[cmp] rows compared: {n} x 21")
print(f"[cmp] max|T_rerun - T_xlsx| = {dT.max():.3e} K at "
      f"(t={int(np.unravel_index(np.argmax(dT), dT.shape)[0])+1} s, "
      f"col={np.unravel_index(np.argmax(dT), dT.shape)[1]})")
print(f"[cmp] max|C_rerun - C_xlsx| = {dC.max():.3e} kg/kg")
print(f"[cmp] T rounded agreement (4 dp): "
      f"{int((np.round(Tr,4)==np.round(Tr_x,4)).sum())}/{dT.size}")
print(f"[cmp] C rounded agreement (4 dp): "
      f"{int((np.round(Cr,4)==np.round(Cr_x,4)).sum())}/{dC.size}")
for ts in (1, 60, 600, 1800):
    i = ts - 1
    print(f"  t={ts:5d} s  T(0): re-run {Tr[i,0]:.4f} xlsx {Tr_x[i,0]:.4f} | "
          f"T(R): {Tr[i,-1]:.4f} / {Tr_x[i,-1]:.4f} | "
          f"C(R): {Cr[i,-1]:.6f} / {Cr_x[i,-1]:.6f}")
    print(f"            full-row max|dT|={dT[i].max():.2e} max|dC|={dC[i].max():.2e}")
