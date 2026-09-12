# -*- coding: utf-8 -*-
"""Audit B: result2.xlsx vs doc tables (C7); extrema/monotone counts (C3); uncertainty (C6)."""
from __future__ import annotations
import io, os, sys
import numpy as np
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "outputs")

def hr(t):
    print("\n" + "=" * 78 + "\n" + t + "\n" + "=" * 78)

hr("C7a  structure of outputs/result2.xlsx vs template")
wb = load_workbook(os.path.join(OUT, "result2.xlsx"), read_only=True, data_only=True)
print("sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"  [{sn}] dims={ws.max_row} x {ws.max_column}")
wsT = wb["温度"]
hdr = [c.value for c in next(wsT.iter_rows(min_row=1, max_row=1))]
print("row1:", hdr)
first = [c.value for c in next(wsT.iter_rows(min_row=2, max_row=2))]
print("row2 (first data row):", first[:5], "...", first[-2:])

# template
tpl = load_workbook(os.path.join(ROOT, "A题", "附件", "附件3", "result2.xlsx"),
                    read_only=True, data_only=True)
print("\nTEMPLATE result2.xlsx")
print("sheets:", tpl.sheetnames)
for sn in tpl.sheetnames:
    ws = tpl[sn]
    print(f"  [{sn}] dims={ws.max_row} x {ws.max_column}")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=2, values_only=True)):
        print(f"    r{i+1}:", row)

# ---- load arrays
def load(sheet):
    ws = wb[sheet]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    tt = np.array([r[0] for r in rows], dtype=float)
    arr = np.array([[np.nan if v is None else float(v) for v in r[1:]] for r in rows])
    return tt, arr
tt, Tc = load("温度")
tt2, Cc = load("水分浓度")
rad = np.array([float(v) for v in hdr[1:]])
print(f"\nloaded T {Tc.shape}, C {Cc.shape}; times {tt[0]:.0f}..{tt[-1]:.0f} "
      f"(n={len(tt)}), radii {rad[0]}..{rad[-1]} (n={len(rad)})")
print("times contiguous 1..10800:", np.array_equal(tt, np.arange(1, 10801)))
print("radii == 0..2 step .1:", np.allclose(rad, np.round(np.arange(0, 2.0001, 0.1), 10)))
print("NaN count:", int(np.isnan(Tc).sum()), int(np.isnan(Cc).sum()))

hr("C7b  every literal in doc section 8.1/8.2 vs result2.xlsx")
T_DOC = [
    [32.1056, 32.2973, 32.8775, 33.8671, 35.3131],
    [40.2754, 40.4472, 40.9536, 41.7705, 42.8919],
    [45.7411, 45.8293, 46.0878, 46.5008, 47.0362],
    [48.3516, 48.3896, 48.5006, 48.6767, 48.9078],
    [49.3772, 49.3896, 49.4246, 49.4765, 49.5751],
    [49.7685, 49.7746, 49.7945, 49.8309, 49.8888]]
C_DOC = [
    [2.5499, 2.5489, 2.5257, 2.3262, 1.6475],
    [2.5261, 2.4953, 2.3587, 2.0235, 1.4699],
    [2.3874, 2.3270, 2.1356, 1.8024, 1.3465],
    [2.1727, 2.1102, 1.9249, 1.6263, 1.2303],
    [1.9586, 1.9024, 1.7373, 1.4725, 1.1161],
    [1.7681, 1.7183, 1.5715, 1.3338, 1.0078]]
T_S = [1800, 3600, 5400, 7200, 9000, 10800]
COL = [0, 5, 10, 15, 20]
bad = 0
for i, ts in enumerate(T_S):
    for j, c in enumerate(COL):
        vT, vC = Tc[ts - 1, c], Cc[ts - 1, c]
        dT, dC = vT - T_DOC[i][j], vC - C_DOC[i][j]
        exactT = abs(dT) < 5e-5
        exactC = abs(dC) < 5e-5
        if not (exactT and exactC):
            bad += 1
            print(f"  MISMATCH t={ts}s r={rad[c]}cm: xlsx T={vT}, doc {T_DOC[i][j]} "
                  f"(d={dT:+.2e}); xlsx C={vC}, doc {C_DOC[i][j]} (d={dC:+.2e})")
print(f"60/60 values agree with the doc table at 4 decimals: {bad == 0}  (mismatches={bad})")

hr("C6  uncertainty arithmetic")
uT, uC = 1.3808e-5, 3.8467e-5
print(f"T: {uT}/5e-5 = {uT/5e-5*100:.4f}%  (doc 27.62%)")
print(f"C: {uC}/5e-5 = {uC/5e-5*100:.4f}%  (doc 76.93%)")
print("dt/2  check: dT=1.3808e-05 dC=2.9839e-05")
print("M x2  check: dT=1.3824e-06 dC=3.8467e-05")
print("max(dt/2, Mx2) T =", max(1.3808e-5, 1.3824e-6), " C =", max(2.9839e-5, 3.8467e-5))

hr("C3  extrema / monotonicity from result2.xlsx")
iT = np.unravel_index(np.argmin(Tc), Tc.shape)
aT = np.unravel_index(np.argmax(Tc), Tc.shape)
print(f"T min {Tc.min():.6f} at t={tt[iT[0]]:.0f}s r={rad[iT[1]]}cm")
print(f"T max {Tc.max():.6f} at t={tt[aT[0]]:.0f}s r={rad[aT[1]]}cm")
print("min over r=2cm col:", Tc[:, -1].min(), "at t=", tt[np.argmin(Tc[:, -1])])
print("min per column (radii):", [f"{rad[j]}:{Tc[:,j].min():.4f}" for j in range(0, 21, 5)])
print("max per column:", [f"{rad[j]}:{Tc[:,j].max():.4f}" for j in range(0, 21, 5)])
iC = np.unravel_index(np.argmin(Cc), Cc.shape)
print(f"C min {Cc.min():.6f} at t={tt[iC[0]]:.0f}s r={rad[iC[1]]}cm ; C max {Cc.max():.6f}")
# time-direction decreases per radius
dec = (np.diff(Tc, axis=0) < 0).sum(axis=0)
print("descending-in-time counts per radius:", dec.tolist())
print("outer (r>=1.5cm) sum:", int(dec[15:].sum()), " inner (r<=1.0cm) sum:", int(dec[:11].sum()))
print("doc table lists r=1.5 ->72, r=1.9 ->632, r=2.0 ->800; actual:",
      dec[15], dec[19], dec[20])
print("C monotone non-increasing per radius (violations):",
      int((np.diff(Cc, axis=0) > 0).sum()))
print("C descending counts per radius:", (np.diff(Cc, axis=0) < 0).sum(axis=0).tolist())

# env noise
sys.path.insert(0, os.path.join(ROOT, "src"))
from q1_solve import load_attachment1
t1, T1C, C1 = load_attachment1()
sel = t1 <= 10800 + 1e-9
ts = t1[sel]; Tv = T1C[sel]
print(f"\nattachment1: n={len(ts)}, t {ts[0]:.0f}..{ts[-1]:.0f}")
d = np.diff(Tv)
print(f"descending env-T steps = {(d < 0).sum()} / {len(d)} ; peak {Tv.max():.4f} at t={ts[np.argmax(Tv)]:.0f}s")
dC1 = np.diff(C1[sel])
print(f"C_inf range [{C1[sel].min():.5f},{C1[sel].max():.5f}]  descending steps {(dC1<0).sum()}")
print("T_inf(0)=", Tv[0], " T_inf(10800)=", Tv[ts == 10800][0] if (ts == 10800).any() else "n/a")
print("doc says C_inf(0)=0.01963, C_inf(10800)=0.04977")

hr("derived numbers in doc section 8.3 / 9")
print("radial span t=10800:", Tc[-1, -1] - Tc[-1, 0])
print("C(0)/C0:", Cc[-1, 0] / 2.55 * 100, "% ; C(R)/C0:", Cc[-1, -1] / 2.55 * 100, "%")
print("T center vs Tinf(10800)=50.1950:", 50.1950 - Tc[-1, 0], "; surface:",
      50.1950 - Tc[-1, -1])
print("doc (b) table: 60s C(0)=", Cc[59, 0], "C(1)=", Cc[59, 10], "C(R)=", Cc[59, 20])
print("doc (b): 600s C(R)=", Cc[599, 20], " (doc says 1.93170)")
